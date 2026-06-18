#!/usr/bin/env python3
"""Build a course-style study deck for this NWM project.

The deck is intentionally self-contained: it explains the problem, repository
architecture, data pipeline, model architecture, training/inference/evaluation
methods, and the current local experiment results.
"""

from __future__ import annotations

from pathlib import Path
import shutil
import subprocess
import textwrap

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "artifacts" / "decks" / "nwm_course"
FIG_DIR = OUT_DIR / "figures"
PPTX_PATH = OUT_DIR / "nwm_project_course.pptx"
PDF_PATH = OUT_DIR / "nwm_project_course.pdf"
NOTES_PATH = OUT_DIR / "nwm_project_course_notes.md"

PAPER_CEM_CSV = (
    ROOT
    / "artifacts"
    / "summaries"
    / "planning"
    / "paper_cem_b_selected_n120_rep3"
    / "metrics.csv"
)
PAPER_CEM_FIG = (
    ROOT
    / "artifacts"
    / "summaries"
    / "planning"
    / "paper_cem_b_selected_n120_rep3"
    / "ate_rpe_comparison.png"
)
TEXT_DELTA_FIG = (
    ROOT
    / "artifacts"
    / "decks"
    / "nwm_text_conditioned_update"
    / "figures"
    / "prediction_text_delta_heatmap.png"
)
PLANNING_GROUPED_FIG = (
    ROOT
    / "artifacts"
    / "decks"
    / "nwm_text_conditioned_update"
    / "figures"
    / "planning_ate_grouped.png"
)

FONT = "Noto Sans CJK KR"
MONO = "DejaVu Sans Mono"

BG = RGBColor(246, 248, 250)
DARK = RGBColor(24, 31, 42)
BODY = RGBColor(52, 61, 73)
MUTED = RGBColor(101, 111, 125)
WHITE = RGBColor(255, 255, 255)
BLUE = RGBColor(37, 99, 162)
TEAL = RGBColor(20, 133, 124)
GREEN = RGBColor(47, 128, 78)
ORANGE = RGBColor(205, 116, 37)
RED = RGBColor(183, 71, 64)
PURPLE = RGBColor(112, 84, 180)
LINE = RGBColor(212, 218, 226)
LIGHT_BLUE = RGBColor(226, 237, 249)
LIGHT_TEAL = RGBColor(224, 244, 241)
LIGHT_GREEN = RGBColor(225, 242, 232)
LIGHT_ORANGE = RGBColor(252, 239, 225)
LIGHT_RED = RGBColor(250, 230, 228)
LIGHT_GRAY = RGBColor(235, 238, 242)

SLIDE_W = 13.333
SLIDE_H = 7.5


class DeckBuilder:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.notes: list[tuple[str, str]] = []
        self.slide_no = 0

    def slide(self, title: str, section: str | None = None, note: str = ""):
        self.slide_no += 1
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = BG

        if section:
            tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(0.32), Inches(1.65), Inches(0.32))
            tag.fill.solid()
            tag.fill.fore_color.rgb = LIGHT_BLUE
            tag.line.color.rgb = LIGHT_BLUE
            tf = tag.text_frame
            tf.clear()
            tf.margin_left = Inches(0.08)
            tf.margin_right = Inches(0.08)
            p = tf.paragraphs[0]
            p.text = section
            p.alignment = PP_ALIGN.CENTER
            p.font.name = FONT
            p.font.size = Pt(9)
            p.font.bold = True
            p.font.color.rgb = BLUE

        self.text(slide, title, 0.55, 0.62, 12.2, 0.45, size=26, bold=True, color=DARK)
        self.line(slide, 0.55, 1.18, 12.2, 0, LINE)
        self.text(slide, f"{self.slide_no:02d}", 12.25, 7.04, 0.55, 0.22, size=8, color=MUTED, align=PP_ALIGN.RIGHT)
        self.notes.append((title, note))
        return slide

    def text(
        self,
        slide,
        text: str,
        x: float,
        y: float,
        w: float,
        h: float,
        size: int = 14,
        bold: bool = False,
        color: RGBColor = BODY,
        align=PP_ALIGN.LEFT,
        font: str = FONT,
        valign=MSO_ANCHOR.TOP,
    ):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.vertical_anchor = valign
        tf.margin_left = Inches(0.02)
        tf.margin_right = Inches(0.02)
        tf.margin_top = Inches(0.02)
        tf.margin_bottom = Inches(0.02)
        p = tf.paragraphs[0]
        p.text = text
        p.alignment = align
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        return box

    def bullets(
        self,
        slide,
        items: list[str],
        x: float,
        y: float,
        w: float,
        h: float,
        size: int = 13,
        color: RGBColor = BODY,
        gap: int = 4,
    ):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.08)
        tf.margin_right = Inches(0.04)
        tf.margin_top = Inches(0.02)
        tf.margin_bottom = Inches(0.02)
        for idx, item in enumerate(items):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = item
            p.level = 0
            p.space_after = Pt(gap)
            p.font.name = FONT
            p.font.size = Pt(size)
            p.font.color.rgb = color
        return box

    def code(self, slide, text: str, x: float, y: float, w: float, h: float, size: int = 8):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(32, 38, 48)
        shape.line.color.rgb = RGBColor(32, 38, 48)
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.10)
        tf.margin_bottom = Inches(0.10)
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = MONO
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(238, 242, 247)
        return shape

    def card(
        self,
        slide,
        title: str,
        body: str,
        x: float,
        y: float,
        w: float,
        h: float,
        fill: RGBColor = WHITE,
        accent: RGBColor = BLUE,
        title_size: int = 13,
        body_size: int = 10,
    ):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = LINE
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.color.rgb = accent
        self.text(slide, title, x + 0.17, y + 0.14, w - 0.32, 0.28, title_size, True, DARK)
        self.text(slide, body, x + 0.17, y + 0.52, w - 0.32, h - 0.62, body_size, False, BODY)
        return shape

    def metric_card(
        self,
        slide,
        label: str,
        value: str,
        caption: str,
        x: float,
        y: float,
        w: float,
        h: float,
        fill: RGBColor,
        accent: RGBColor,
    ):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.color.rgb = fill
        self.text(slide, label, x + 0.15, y + 0.16, w - 0.3, 0.25, 10, True, accent)
        self.text(slide, value, x + 0.15, y + 0.48, w - 0.3, 0.50, 22, True, DARK)
        self.text(slide, caption, x + 0.15, y + 1.04, w - 0.3, h - 1.05, 9, False, BODY)
        return shape

    def line(self, slide, x: float, y: float, w: float, h: float, color: RGBColor):
        line = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y + h))
        line.line.color.rgb = color
        line.line.width = Pt(1)
        return line

    def flow(self, slide, steps: list[tuple[str, str]], x: float, y: float, w: float, h: float, colors: list[RGBColor] | None = None):
        n = len(steps)
        gap = 0.18
        arrow_w = 0.22
        box_w = (w - (n - 1) * (gap + arrow_w)) / n
        for i, (title, body) in enumerate(steps):
            bx = x + i * (box_w + gap + arrow_w)
            fill = colors[i] if colors else WHITE
            self.card(slide, title, body, bx, y, box_w, h, fill=fill, accent=[BLUE, TEAL, ORANGE, GREEN, PURPLE][i % 5], title_size=10, body_size=8)
            if i < n - 1:
                arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(bx + box_w + 0.04), Inches(y + h / 2 - 0.11), Inches(arrow_w), Inches(0.22))
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = MUTED
                arrow.line.color.rgb = MUTED

    def table(
        self,
        slide,
        headers: list[str],
        rows: list[list[str]],
        x: float,
        y: float,
        w: float,
        h: float,
        font_size: int = 8,
    ):
        shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h))
        table = shape.table
        for j, header in enumerate(headers):
            cell = table.cell(0, j)
            cell.text = header
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT_BLUE
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(font_size)
                p.font.bold = True
                p.font.color.rgb = DARK
                p.alignment = PP_ALIGN.CENTER
        for i, row in enumerate(rows, start=1):
            for j, value in enumerate(row):
                cell = table.cell(i, j)
                cell.text = value
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if i % 2 else RGBColor(249, 250, 252)
                for p in cell.text_frame.paragraphs:
                    p.font.name = FONT
                    p.font.size = Pt(font_size)
                    p.font.color.rgb = BODY
                    p.alignment = PP_ALIGN.CENTER if j > 0 else PP_ALIGN.LEFT
        return shape

    def image(self, slide, path: Path, x: float, y: float, w: float, h: float | None = None):
        if path.exists():
            if h is None:
                return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
            return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
        self.card(slide, "Figure missing", str(path.relative_to(ROOT)), x, y, w, h or 1.0, fill=LIGHT_RED, accent=RED)
        return None


def ensure_dirs() -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)


def load_cem_results() -> pd.DataFrame:
    if PAPER_CEM_CSV.exists():
        df = pd.read_csv(PAPER_CEM_CSV)
        rename = {
            "label": "variant",
            "rpe_trans": "rpe",
            "pos_error": "pos",
            "yaw_error": "yaw",
            "ate_delta_vs_n32_pct": "ate_vs_n32_pct",
            "rpe_trans_delta_vs_n32_pct": "rpe_vs_n32_pct",
            "ate_delta_vs_paper_pct": "ate_vs_paper_pct",
            "rpe_delta_vs_paper_pct": "rpe_vs_paper_pct",
        }
        return df.rename(columns={k: v for k, v in rename.items() if k in df.columns})
    return pd.DataFrame(
        [
            ["CDiT-B 224 No-Text", "N120/K5/rep3/OPT1", 100, "done", 1.130, 0.348, 1.627, 0.265, 52342.6, -2.72, -2.09, -0.01, -0.57],
            ["CDiT-B 224 Text", "N120/K5/rep3/OPT1", 100, "done", 1.293, 0.350, 1.722, 0.154, 52415.7, 2.56, -1.17, 14.42, 0.00],
            ["CDiT-B 128 Text", "N120/K5/rep3/OPT1", 100, "done", 1.244, 0.344, 1.718, 0.182, 34230.2, -1.28, -3.99, 10.12, -1.71],
        ],
        columns=["variant", "setting", "samples", "status", "ate", "rpe", "pos", "yaw", "time_s", "ate_vs_n32_pct", "rpe_vs_n32_pct", "ate_vs_paper_pct", "rpe_vs_paper_pct"],
    )


def build_cem_chart(df: pd.DataFrame) -> Path:
    out = FIG_DIR / "paper_cem_course_chart.png"
    full = df[(df["setting"] == "N120/K5/rep3/OPT1") & (df["status"] == "done")].copy()
    if full.empty:
        return out
    labels = full["variant"].str.replace("CDiT-B ", "", regex=False).tolist()
    x = np.arange(len(full))
    fig, axes = plt.subplots(1, 2, figsize=(10, 3.8))
    colors = ["#2f80c0", "#d07a2c", "#198f7a"]
    axes[0].bar(x, full["ate"], color=colors[: len(full)])
    axes[0].axhline(1.13, color="#7b8490", linestyle="--", linewidth=1.2, label="Paper NWM ATE 1.13")
    axes[0].set_title("ATE (lower is better)")
    axes[0].set_xticks(x, labels, rotation=18, ha="right")
    axes[0].set_ylim(0, max(1.42, full["ate"].max() * 1.15))
    axes[0].legend(fontsize=7, frameon=False)
    for i, value in enumerate(full["ate"]):
        axes[0].text(i, value + 0.025, f"{value:.3f}", ha="center", fontsize=8)

    axes[1].bar(x, full["rpe"], color=colors[: len(full)])
    axes[1].axhline(0.35, color="#7b8490", linestyle="--", linewidth=1.2, label="Paper NWM RPE 0.35")
    axes[1].set_title("RPE trans (lower is better)")
    axes[1].set_xticks(x, labels, rotation=18, ha="right")
    axes[1].set_ylim(0, max(0.40, full["rpe"].max() * 1.18))
    axes[1].legend(fontsize=7, frameon=False)
    for i, value in enumerate(full["rpe"]):
        axes[1].text(i, value + 0.007, f"{value:.3f}", ha="center", fontsize=8)

    for ax in axes:
        ax.grid(axis="y", alpha=0.25)
        ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(out, bbox_inches="tight")
    plt.close(fig)
    return out


def key_cem_rows(df: pd.DataFrame) -> list[list[str]]:
    full = df[(df["setting"] == "N120/K5/rep3/OPT1") & (df["status"] == "done")].copy()
    rows = []
    for _, r in full.iterrows():
        rows.append(
            [
                str(r["variant"]).replace("CDiT-B ", "B "),
                f"{float(r['ate']):.3f}",
                f"{float(r['rpe']):.3f}",
                f"{float(r['ate_vs_n32_pct']):+.2f}%",
                f"{float(r['rpe_vs_n32_pct']):+.2f}%",
                f"{float(r['ate_vs_paper_pct']):+.2f}%",
            ]
        )
    return rows


def maybe_export_pdf() -> None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return
    subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(OUT_DIR),
            str(PPTX_PATH),
        ],
        check=False,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )


def build_deck() -> None:
    ensure_dirs()
    cem_df = load_cem_results()
    cem_chart = build_cem_chart(cem_df)

    deck = DeckBuilder()

    # 1
    slide = deck.slide(
        "Navigation World Models 프로젝트 강의",
        "0. 안내",
        "프로젝트 전체를 한 과목처럼 학습할 수 있도록 문제, 데이터, 모델, 학습, 평가, 실험 결과를 순서대로 정리한 덱.",
    )
    deck.text(slide, "Text-Conditioned Lightweight NWM / RECON Fork", 0.7, 1.45, 8.2, 0.45, size=18, bold=True, color=BLUE)
    deck.text(slide, "이 PPT의 목표는 코드를 바로 실행하기 전에 전체 시스템의 목적과 흐름을 먼저 머릿속에 그리는 것입니다.", 0.7, 2.02, 8.8, 0.62, size=14, color=BODY)
    deck.metric_card(slide, "Course Mode", "31 slides", "개념 → 코드 → 실험 → 실습 순서", 0.78, 3.05, 2.55, 1.55, LIGHT_BLUE, BLUE)
    deck.metric_card(slide, "Core Model", "CDiT", "Conditional Diffusion Transformer", 3.55, 3.05, 2.55, 1.55, LIGHT_TEAL, TEAL)
    deck.metric_card(slide, "Main Dataset", "RECON", "Trajectory frames + pose/action", 6.32, 3.05, 2.55, 1.55, LIGHT_ORANGE, ORANGE)
    deck.metric_card(slide, "Current Variant", "Text", "Qwen captions + CLIP embeddings", 9.09, 3.05, 2.75, 1.55, LIGHT_GREEN, GREEN)
    deck.code(slide, "repo: /home/kun/kun_ssd/nwm\noutputs: artifacts/decks/nwm_course/", 0.78, 5.45, 11.8, 0.72, size=10)

    # 2
    slide = deck.slide(
        "수업 로드맵",
        "0. 안내",
        "전체 학습 순서와 각 파트에서 얻어야 하는 직관.",
    )
    deck.flow(
        slide,
        [
            ("문제", "미래 시각 상태를 예측해 navigation을 돕는다."),
            ("데이터", "trajectory, 이미지, pose, action, text cache를 다룬다."),
            ("모델", "VAE latent + diffusion + CDiT conditioning."),
            ("평가", "time, rollout, CEM planning으로 검증한다."),
            ("실습", "Docker, smoke, train, infer, evaluate 명령."),
        ],
        0.75,
        1.65,
        11.8,
        1.55,
        [WHITE, WHITE, WHITE, WHITE, WHITE],
    )
    deck.card(slide, "학습 목표 1", "이 repo에서 어떤 파일이 무엇을 담당하는지 설명할 수 있다.", 0.85, 4.05, 3.55, 1.15, fill=WHITE, accent=BLUE)
    deck.card(slide, "학습 목표 2", "CDiT world model이 무엇을 예측하고 어떤 조건을 쓰는지 설명할 수 있다.", 4.85, 4.05, 3.55, 1.15, fill=WHITE, accent=TEAL)
    deck.card(slide, "학습 목표 3", "평가 결과에서 LPIPS/DreamSim/FID/ATE/RPE가 의미하는 바를 읽을 수 있다.", 8.85, 4.05, 3.55, 1.15, fill=WHITE, accent=ORANGE)

    # 3
    slide = deck.slide(
        "프로젝트 한 줄 정의",
        "1. 문제",
        "NWM이 풀고 싶은 문제를 정책 학습이 아니라 world modeling 관점에서 정의.",
    )
    deck.text(slide, "Navigation World Model은 현재 관측과 행동 후보를 조건으로 가까운 미래의 시각 상태를 생성하는 모델입니다.", 0.75, 1.55, 11.8, 0.65, 20, True, DARK)
    deck.card(slide, "입력", "과거 context frames\n상대 이동/회전 action\n미래 시간 offset\n선택적 text embedding", 0.85, 2.65, 2.6, 2.2, fill=LIGHT_BLUE, accent=BLUE)
    deck.card(slide, "출력", "미래 프레임 또는 rollout\nVAE latent를 거쳐 pixel image로 복원\n후속 평가/계획에 사용", 3.75, 2.65, 2.6, 2.2, fill=LIGHT_TEAL, accent=TEAL)
    deck.card(slide, "학습", "정책을 직접 학습하지 않음\n행동 조건부 미래 latent의 denoising을 학습", 6.65, 2.65, 2.6, 2.2, fill=LIGHT_ORANGE, accent=ORANGE)
    deck.card(slide, "계획", "후보 action을 샘플링하고\nworld model rollout을 goal image와 비교해 고름", 9.55, 2.65, 2.6, 2.2, fill=LIGHT_GREEN, accent=GREEN)
    deck.text(slide, "핵심 질문: “이 행동을 하면 몇 초 뒤 내 카메라가 무엇을 볼까?”", 0.95, 5.65, 11.1, 0.45, 17, True, BLUE, align=PP_ALIGN.CENTER)

    # 4
    slide = deck.slide(
        "이 프로젝트에서 우리가 추가로 묻는 질문",
        "1. 문제",
        "원본 NWM 위에 text conditioning을 얹은 이 fork의 연구 질문.",
    )
    deck.bullets(
        slide,
        [
            "저해상도 또는 경량 모델에서는 이미지에서 길, 벽, 통로, 장애물의 의미가 흐려질 수 있다.",
            "VLM caption을 offline으로 만들고 CLIP embedding으로 cache하면 semantic prior로 쓸 수 있다.",
            "텍스트는 명령이 아니라 장면의 의미를 보완하는 조건 신호다.",
            "검증은 frame prediction, rollout, planning 성능으로 나눠서 본다.",
        ],
        0.85,
        1.55,
        6.25,
        2.5,
        size=14,
    )
    deck.flow(
        slide,
        [
            ("Raw frame", "RECON image"),
            ("Qwen caption", "offline scene text"),
            ("CLIP embed", "dense npz cache"),
            ("CDiT cond", "text_proj로 합산"),
        ],
        0.85,
        4.55,
        6.25,
        1.2,
        [WHITE, WHITE, WHITE, WHITE],
    )
    deck.card(slide, "주의", "텍스트 조건을 넣는다고 planning policy를 학습하는 것은 아닙니다. 모델 학습은 여전히 future image/latent prediction이고, CEM은 평가 시 후보 행동을 고르는 planner입니다.", 7.55, 1.55, 4.75, 4.2, fill=WHITE, accent=RED, title_size=15, body_size=12)

    # 5
    slide = deck.slide(
        "Repository Architecture",
        "2. 코드 구조",
        "폴더 단위로 이 repo가 어떻게 나뉘는지.",
    )
    deck.code(
        slide,
        """nwm/
├── scripts/       실행 진입점: train, infer, evaluate, plan_eval
├── src/           실제 Python package
│   ├── data/      dataset, trajectory/action/image transform
│   ├── models/    CDiT backbone, checkpoint, VAE loader
│   ├── diffusion/ DDPM / respacing
│   ├── evaluation/inference, metrics, planning
│   └── features/text/ text cache utilities
├── configs/       experiment/data/evaluation YAML
├── data/splits/   train/test/eval split metadata
├── datasets/      raw and derived datasets
├── weights/       checkpoints, pretrained, cache
├── logs/          train/runtime logs
└── artifacts/     eval outputs, summaries, decks, figures""",
        0.75,
        1.45,
        6.1,
        5.45,
        size=9,
    )
    deck.card(slide, "읽는 순서", "1. README / docs/project_structure\n2. configs/experiment/*.yaml\n3. scripts/train.py\n4. src/models/backbones/cdit.py\n5. scripts/infer.py / evaluate.py / plan_eval.py", 7.2, 1.52, 4.9, 2.15, fill=WHITE, accent=BLUE)
    deck.card(slide, "중요한 구분", "실제 모델 코드는 `src/models/backbones/cdit.py`입니다. diffusion 구현은 `src/diffusion/`이고, pretrained weight는 `weights/pretrained/` 아래에 둡니다.", 7.2, 4.05, 4.9, 1.55, fill=LIGHT_ORANGE, accent=ORANGE)

    # 6
    slide = deck.slide(
        "Runtime Architecture",
        "2. 코드 구조",
        "config, entrypoint, library module, artifact가 이어지는 실행 구조.",
    )
    deck.flow(
        slide,
        [
            ("YAML config", "model, dataset, image_size, text"),
            ("script", "train.py / infer.py / plan_eval.py"),
            ("src modules", "data + model + diffusion + metrics"),
            ("outputs", "weights + logs + artifacts"),
        ],
        0.75,
        1.55,
        11.6,
        1.55,
        [LIGHT_BLUE, LIGHT_TEAL, LIGHT_ORANGE, LIGHT_GREEN],
    )
    deck.card(slide, "Train path", "`configs/experiment/*.yaml` → `TrainingDataset` → `CDiT` → `create_diffusion` → `weights/checkpoints/<run_name>`", 0.85, 3.75, 3.6, 1.35, fill=WHITE, accent=BLUE)
    deck.card(slide, "Prediction path", "`infer.py` loads EMA checkpoint → generates GT or predictions → writes frames under eval artifact root", 4.85, 3.75, 3.6, 1.35, fill=WHITE, accent=TEAL)
    deck.card(slide, "Planning path", "`plan_eval.py` → CEM action sampling → world model rollout → ATE/RPE metrics JSON", 8.85, 3.75, 3.6, 1.35, fill=WHITE, accent=ORANGE)
    deck.text(slide, "핵심 패턴: config는 실험 정의, scripts는 실행 버튼, src는 재사용 가능한 구현입니다.", 0.95, 6.0, 11.2, 0.35, 14, True, DARK, align=PP_ALIGN.CENTER)

    # 7
    slide = deck.slide(
        "데이터 단위: trajectory sample",
        "3. 데이터",
        "TrainingDataset이 한 샘플에서 무엇을 만드는지.",
    )
    deck.flow(
        slide,
        [
            ("trajectory", "0.jpg ... T.jpg\ntraj_data.pkl"),
            ("context", "curr_time 이전\ncontext_size frames"),
            ("goal", "future offset\nmin/max distance"),
            ("action", "local x, y, yaw\nnormalized"),
            ("text", "current/goal/context_mean\noptional"),
        ],
        0.65,
        1.55,
        12.05,
        1.55,
        [WHITE, WHITE, WHITE, WHITE, WHITE],
    )
    deck.card(slide, "학습 샘플", "`obs_image`: context frames + goal frames\n`goal_pos`: local coordinate action target\n`rel_time`: goal offset / 128\n`text_emb`: optional CLIP vector", 0.85, 3.65, 3.8, 2.05, fill=LIGHT_BLUE, accent=BLUE)
    deck.card(slide, "데이터셋 코드", "`src/data/datasets/base_dataset.py`\nindex 생성, trajectory 로드, action 계산, text cache lookup\n\n`src/data/datasets/train_dataset.py`\nrandom goal offset sampling", 4.85, 3.65, 3.8, 2.05, fill=LIGHT_TEAL, accent=TEAL)
    deck.card(slide, "평가 데이터", "time/rollout/planning은 predefined index를 사용해 같은 샘플 순서로 비교합니다.\n예: `data/splits/recon/test/navigation_eval.pkl`", 8.85, 3.65, 3.8, 2.05, fill=LIGHT_ORANGE, accent=ORANGE)

    # 8
    slide = deck.slide(
        "RECON + Text Pipeline",
        "3. 데이터",
        "이 fork에서 text-conditioned 실험을 위해 추가된 offline pipeline.",
    )
    deck.flow(
        slide,
        [
            ("Raw RECON", "HDF5 / frames\npose"),
            ("1fps export", "caption 대상\nframe subset"),
            ("Qwen2-VL", "scene caption"),
            ("cleaning", "boilerplate 제거"),
            ("CLIP", "text embedding"),
            ("dense cache", "trajectory .npz"),
        ],
        0.55,
        1.55,
        12.25,
        1.25,
        [WHITE, WHITE, WHITE, WHITE, WHITE, WHITE],
    )
    deck.bullets(
        slide,
        [
            "`scripts/preprocess/recon/`: RECON export/render helpers",
            "`scripts/preprocess/text/`: Qwen caption, cleaning, CLIP embedding, dense alignment",
            "`datasets/derived/phase1_text_embeds_dense/recon_all_raw_rel`: 학습 시 읽는 dense text cache",
            "`text_conditioning.enabled: True`이면 dataset loader가 text embedding을 batch에 추가",
        ],
        0.85,
        3.45,
        7.3,
        2.1,
        size=13,
    )
    deck.card(slide, "왜 offline인가?", "학습/평가 루프 안에서 VLM을 돌리면 너무 느립니다. caption과 embedding을 미리 만들어 cache하면 기존 NWM 실행 속도와 구조를 크게 깨지 않고 semantic signal을 추가할 수 있습니다.", 8.65, 3.35, 3.75, 2.25, fill=WHITE, accent=GREEN)

    # 9
    slide = deck.slide(
        "무엇을 학습하는가",
        "4. 모델",
        "정책 학습이 아니라 조건부 미래 latent denoising을 학습한다는 점을 명확히 함.",
    )
    deck.text(slide, "학습 목표", 0.85, 1.45, 2.5, 0.3, 17, True, BLUE)
    deck.code(
        slide,
        "given: context frames, action(x,y,yaw), relative time, optional text\n"
        "target: future frame latent z_goal\n"
        "diffusion: add noise to z_goal at timestep t\n"
        "model: predict denoising target conditioned on inputs\n\n"
        "learn p(z_future | z_context, action, rel_time, text)",
        0.85,
        1.95,
        5.9,
        2.25,
        size=11,
    )
    deck.card(slide, "모델이 직접 배우는 것", "미래 이미지 latent의 denoising\n관측과 행동 사이의 시각적 dynamics\n시간 offset에 따른 변화", 7.25, 1.75, 4.6, 1.45, fill=LIGHT_TEAL, accent=TEAL)
    deck.card(slide, "모델이 직접 배우지 않는 것", "최적 policy\nreward function\n장기 안전성 보장\nCEM hyperparameter 자체", 7.25, 3.45, 4.6, 1.45, fill=LIGHT_RED, accent=RED)
    deck.text(slide, "즉, world model은 “행동의 결과를 상상하는 모델”이고 planner는 그 상상을 이용해 행동을 고릅니다.", 0.95, 5.75, 11.2, 0.5, 15, True, DARK, align=PP_ALIGN.CENTER)

    # 10
    slide = deck.slide(
        "Latent Diffusion 구조",
        "4. 모델",
        "이미지를 직접 diffusion하지 않고 VAE latent에서 학습하는 이유와 흐름.",
    )
    deck.flow(
        slide,
        [
            ("Image", "RGB frame"),
            ("VAE encoder", "latent z\nx / 8"),
            ("DDPM noise", "z_t"),
            ("CDiT", "denoise"),
            ("VAE decoder", "pred frame"),
        ],
        0.8,
        1.75,
        11.6,
        1.55,
        [WHITE, WHITE, WHITE, WHITE, WHITE],
    )
    deck.card(slide, "VAE", "`load_vae()`로 StabilityAI VAE를 로드합니다.\n이미지는 latent로 인코딩되고 `0.18215` scaling을 적용합니다.", 0.9, 4.05, 3.55, 1.55, fill=LIGHT_BLUE, accent=BLUE)
    deck.card(slide, "Diffusion", "`create_diffusion()`은 학습 시 1000 step DDPM, 추론/평가에서는 respacing된 sampling step을 사용합니다.", 4.9, 4.05, 3.55, 1.55, fill=LIGHT_TEAL, accent=TEAL)
    deck.card(slide, "장점", "픽셀 공간보다 latent 공간이 작아서 학습/추론 비용이 줄고, CDiT가 spatial patch token을 다루기 쉬워집니다.", 8.9, 4.05, 3.55, 1.55, fill=LIGHT_GREEN, accent=GREEN)

    # 11
    slide = deck.slide(
        "CDiT 전체 아키텍처",
        "4. 모델",
        "src/models/backbones/cdit.py 기준 모델 구조.",
    )
    deck.flow(
        slide,
        [
            ("Noisy target", "future latent\npatch tokens"),
            ("Context", "past latent tokens\ncross-attn memory"),
            ("Condition c", "t + action + rel_t + text"),
            ("CDiT blocks", "self-attn\ncross-attn\nMLP"),
            ("Final layer", "pred noise/sigma\nunpatchify"),
        ],
        0.55,
        1.5,
        12.25,
        1.55,
        [LIGHT_BLUE, LIGHT_TEAL, LIGHT_ORANGE, LIGHT_GREEN, WHITE],
    )
    deck.card(slide, "Model variants", "CDiT-S/2: 12 layers, 384 hidden, 6 heads\nCDiT-B/2: 12 layers, 768 hidden, 12 heads\nCDiT-L/2: 24 layers, 1024 hidden, 16 heads\nCDiT-XL/2: 28 layers, 1152 hidden, 16 heads", 0.85, 3.75, 4.1, 2.05, fill=WHITE, accent=BLUE)
    deck.card(slide, "핵심 conditioning", "`TimestepEmbedder`: diffusion timestep\n`ActionEmbedder`: x/y/yaw를 각각 embedding\n`time_embedder`: relative future offset\n`text_proj`: CLIP text vector → hidden size", 5.2, 3.75, 3.55, 2.05, fill=WHITE, accent=TEAL)
    deck.card(slide, "출력", "`learn_sigma=True`이면 output channel은 latent channel의 2배입니다.\nfinal layer 후 unpatchify하여 diffusion loss에 들어갈 prediction을 만듭니다.", 9.0, 3.75, 3.55, 2.05, fill=WHITE, accent=ORANGE)

    # 12
    slide = deck.slide(
        "CDiTBlock 내부",
        "4. 모델",
        "한 블록에서 target token과 context token이 어떻게 상호작용하는지.",
    )
    deck.card(slide, "1. Self-Attention", "noisy target patch token끼리 spatial relation을 봅니다.", 0.8, 1.65, 2.75, 1.2, fill=LIGHT_BLUE, accent=BLUE)
    deck.card(slide, "2. Cross-Attention", "target token이 context frame token을 key/value로 참조합니다.", 4.0, 1.65, 2.75, 1.2, fill=LIGHT_TEAL, accent=TEAL)
    deck.card(slide, "3. MLP", "token별 feature를 비선형 변환합니다.", 7.2, 1.65, 2.75, 1.2, fill=LIGHT_ORANGE, accent=ORANGE)
    deck.card(slide, "4. adaLN-Zero", "condition vector c가 각 sub-layer의 shift/scale/gate를 만듭니다.", 10.4, 1.65, 2.0, 1.2, fill=LIGHT_GREEN, accent=GREEN)
    deck.line(slide, 3.55, 2.25, 0.42, 0, MUTED)
    deck.line(slide, 6.75, 2.25, 0.42, 0, MUTED)
    deck.line(slide, 9.95, 2.25, 0.42, 0, MUTED)
    deck.text(slide, "block input", 0.95, 4.05, 1.4, 0.3, 11, True, MUTED)
    deck.code(slide, "x      = noisy target tokens\nx_cond = context frame tokens\nc      = t_embed + action_embed + rel_time_embed (+ text_proj)", 0.95, 4.45, 5.45, 1.1, size=10)
    deck.text(slide, "block output", 7.15, 4.05, 1.6, 0.3, 11, True, MUTED)
    deck.code(slide, "x = x + gated self-attn\nx = x + gated cross-attn(x_cond)\nx = x + gated MLP", 7.15, 4.45, 4.65, 1.1, size=10)

    # 13
    slide = deck.slide(
        "Text Conditioning은 어디에 들어가나",
        "4. 모델",
        "텍스트가 이미지 token에 직접 붙는 것이 아니라 condition vector에 더해짐.",
    )
    deck.flow(
        slide,
        [
            ("caption", "visible scene"),
            ("CLIP text encoder", "embedding"),
            ("text_proj", "hidden size"),
            ("condition c", "t + action + rel_t + text"),
            ("adaLN", "all CDiT blocks"),
        ],
        0.75,
        1.55,
        11.7,
        1.55,
        [WHITE, WHITE, WHITE, WHITE, WHITE],
    )
    deck.card(slide, "선택 가능한 source", "`current`: 현재 프레임 caption\n`goal`: goal frame caption\n`context_mean`: context caption 평균", 0.85, 3.75, 3.45, 1.6, fill=LIGHT_BLUE, accent=BLUE)
    deck.card(slide, "현재 주요 설정", "`text_conditioning.enabled: True`\n`embedding_root: datasets/derived/.../recon_all_raw_rel`\n`condition_source: current`", 4.85, 3.75, 3.45, 1.6, fill=LIGHT_TEAL, accent=TEAL)
    deck.card(slide, "해석", "텍스트는 token sequence로 cross-attention되는 방식이 아니라, diffusion denoising의 global condition을 보정하는 semantic prior입니다.", 8.85, 3.75, 3.45, 1.6, fill=LIGHT_ORANGE, accent=ORANGE)

    # 14
    slide = deck.slide(
        "Training Loop",
        "5. 학습",
        "scripts/train.py의 실제 학습 흐름.",
    )
    deck.flow(
        slide,
        [
            ("batch", "context+goal\nframes"),
            ("VAE encode", "latent tensor"),
            ("sample t", "diffusion step"),
            ("CDiT loss", "training_losses"),
            ("optimizer", "AdamW"),
            ("EMA", "eval copy"),
        ],
        0.55,
        1.55,
        12.25,
        1.35,
        [WHITE, WHITE, WHITE, WHITE, WHITE, WHITE],
    )
    deck.bullets(
        slide,
        [
            "DDP 환경에서 rank별 DataLoader를 돌리고, `DistributedSampler`로 shuffle을 맞춥니다.",
            "context frame은 conditioning latent가 되고, goal frame latent는 diffusion target이 됩니다.",
            "`diffusion.training_losses(model, x_start, t, model_kwargs)`가 핵심 loss를 계산합니다.",
            "EMA 모델은 checkpoint에 함께 저장되고, inference/evaluation에서 EMA weight를 사용합니다.",
            "선택적으로 bfloat16, GradScaler, torch.compile, WandB logging을 사용합니다.",
        ],
        0.9,
        3.35,
        11.25,
        2.25,
        size=13,
    )
    deck.code(slide, "python scripts/train.py --config configs/experiment/nwm_cdit_b_recon_raw_text_dense.yaml --ckpt-every 2000 --eval-every 10000", 0.9, 6.05, 11.3, 0.45, size=9)

    # 15
    slide = deck.slide(
        "Config를 읽는 법",
        "5. 학습",
        "실험 YAML에서 가장 먼저 봐야 할 key.",
    )
    deck.code(
        slide,
        """run_name: nwm_cdit_b_recon_raw_text_dense
model: CDiT-B/2
image_size: 224
context_size: 4
len_traj_pred: 64
lr: 8e-5
from_checkpoint: weights/checkpoints/nwm_cdit_b/0100000.pth.tar
text_conditioning:
  enabled: True
  embedding_root: datasets/derived/phase1_text_embeds_dense/recon_all_raw_rel
  condition_source: current
datasets:
  recon:
    data_folder: datasets/recon_raw/recon_release
    train: data/splits/recon/train/
    test: data/splits/recon/test/""",
        0.8,
        1.55,
        6.05,
        4.7,
        size=8,
    )
    deck.card(slide, "해석 포인트", "`model`과 `image_size`가 architecture와 latent grid size를 결정합니다.\n\n`from_checkpoint`는 warm-start 여부입니다.\n\n`text_conditioning` block이 있으면 model 생성 시 `text_dim`이 들어가고 dataset이 text cache를 로드합니다.", 7.25, 1.75, 4.75, 3.2, fill=WHITE, accent=BLUE)

    # 16
    slide = deck.slide(
        "Inference: time prediction vs rollout",
        "6. 추론",
        "scripts/infer.py가 생성하는 두 가지 평가용 prediction.",
    )
    deck.card(slide, "Time prediction", "1s, 2s, 4s, 8s, 16s 같은 특정 미래 offset의 frame을 한 번에 예측합니다.\n\nGT와 직접 비교하기 쉽고 horizon별 품질을 봅니다.", 0.9, 1.65, 5.2, 2.25, fill=LIGHT_BLUE, accent=BLUE)
    deck.card(slide, "Autoregressive rollout", "한 step 예측 결과를 다음 context로 다시 넣어 trajectory를 이어갑니다.\n\n누적 오류, drift, 장기 안정성을 보기에 좋습니다.", 7.0, 1.65, 5.2, 2.25, fill=LIGHT_TEAL, accent=TEAL)
    deck.flow(
        slide,
        [
            ("context", "obs frames"),
            ("model", "next frame"),
            ("append", "new context"),
            ("repeat", "rollout"),
        ],
        1.2,
        4.65,
        10.9,
        1.2,
        [WHITE, WHITE, WHITE, WHITE],
    )

    # 17
    slide = deck.slide(
        "Prediction Metrics",
        "7. 평가",
        "이미지 생성 품질을 평가하는 metric.",
    )
    deck.card(slide, "LPIPS", "VGG/AlexNet feature 공간에서 perceptual distance를 계산합니다. 낮을수록 GT와 지각적으로 비슷합니다.", 0.85, 1.65, 3.6, 1.65, fill=LIGHT_BLUE, accent=BLUE)
    deck.card(slide, "DreamSim", "인간 유사도 판단에 맞춘 feature 기반 거리입니다. 낮을수록 좋습니다.", 4.85, 1.65, 3.6, 1.65, fill=LIGHT_TEAL, accent=TEAL)
    deck.card(slide, "FID", "예측 이미지 분포와 GT 이미지 분포의 feature statistics 차이를 봅니다. 낮을수록 좋습니다.", 8.85, 1.65, 3.6, 1.65, fill=LIGHT_ORANGE, accent=ORANGE)
    deck.bullets(
        slide,
        [
            "`infer.py --gt 1`로 GT frame set을 먼저 저장합니다.",
            "`infer.py` prediction mode로 모델 결과를 같은 구조에 저장합니다.",
            "`evaluate.py`가 GT directory와 experiment directory를 비교해 JSON metric을 씁니다.",
        ],
        1.0,
        4.2,
        10.9,
        1.3,
        size=13,
    )
    deck.code(slide, "python scripts/evaluate.py --datasets recon --gt_dir <gt> --exp_dir <pred> --eval_types time,rollout", 1.0, 5.95, 10.9, 0.42, size=9)

    # 18
    slide = deck.slide(
        "Planning Evaluation: CEM",
        "7. 평가",
        "world model을 사용해 action 후보를 고르는 평가 방식.",
    )
    deck.flow(
        slide,
        [
            ("sample actions", "mu/sigma에서 N개"),
            ("rollout", "world model 예측"),
            ("score", "goal image LPIPS"),
            ("top-k", "좋은 후보 선택"),
            ("refit", "mu/sigma 갱신"),
            ("metrics", "ATE/RPE"),
        ],
        0.55,
        1.55,
        12.25,
        1.35,
        [WHITE, WHITE, WHITE, WHITE, WHITE, WHITE],
    )
    deck.card(slide, "CEM은 평가 방법", "CEM hyperparameter를 바꾸는 것은 모델 재학습이 아닙니다. checkpoint는 그대로 두고, evaluation/planning stage에서 후보 action 탐색 강도를 바꾸는 것입니다.", 0.85, 3.45, 5.1, 1.85, fill=LIGHT_RED, accent=RED)
    deck.card(slide, "Paper-style local setting", "`N120/K5/rep3/OPT1`\n`num_samples=120`\n`topk=5`\n`num_repeat_eval=3`\n`opt_steps=1`\n`rollout_stride=1`\n`batch_size=1`", 6.55, 3.45, 5.1, 1.85, fill=LIGHT_BLUE, accent=BLUE)

    # 19
    slide = deck.slide(
        "Trajectory Metrics",
        "7. 평가",
        "Planning 결과는 이미지 품질이 아니라 trajectory error로 봅니다.",
    )
    deck.card(slide, "ATE", "Absolute Trajectory Error\n전체 predicted trajectory가 GT trajectory와 얼마나 떨어져 있는지 봅니다. 낮을수록 좋습니다.", 0.9, 1.65, 3.55, 1.8, fill=LIGHT_BLUE, accent=BLUE)
    deck.card(slide, "RPE", "Relative Pose Error\n인접 pose 변화의 상대 오차를 봅니다. local motion consistency를 반영합니다.", 4.9, 1.65, 3.55, 1.8, fill=LIGHT_TEAL, accent=TEAL)
    deck.card(slide, "Pos/Yaw", "최종 위치 error와 yaw error입니다. 실제 navigation에서 goal 근처로 가는지를 보완적으로 확인합니다.", 8.9, 1.65, 3.55, 1.8, fill=LIGHT_ORANGE, accent=ORANGE)
    deck.text(slide, "Planning metric은 prediction metric과 다릅니다. 이미지가 그럴듯해도 action trajectory가 틀릴 수 있고, 반대로 일부 이미지 artifact가 있어도 trajectory가 맞을 수 있습니다.", 1.1, 4.45, 11.0, 0.85, 15, True, DARK, align=PP_ALIGN.CENTER)

    # 20
    slide = deck.slide(
        "현재 Paper-Style CEM 결과",
        "8. 결과",
        "방금 완료된 N120/K5/rep3/OPT1 full 100-sample 결과.",
    )
    deck.image(slide, cem_chart, 0.85, 1.42, 6.6, 2.65)
    deck.table(
        slide,
        ["Variant", "ATE", "RPE", "ATE vs N32", "RPE vs N32", "ATE vs Paper"],
        key_cem_rows(cem_df),
        0.85,
        4.35,
        11.65,
        1.45,
        font_size=8,
    )
    deck.card(slide, "읽는 법", "B224 No-Text는 논문 NWM-only와 거의 동률입니다. B224 Text는 RPE는 거의 동률/소폭 개선이지만 ATE는 N32와 논문 기준보다 나쁩니다. B128 Text는 N32 대비 ATE/RPE 모두 개선이나 논문 ATE에는 못 미칩니다.", 7.85, 1.55, 4.25, 2.35, fill=WHITE, accent=BLUE, body_size=10)

    # 21
    slide = deck.slide(
        "Frame Prediction 결과를 읽는 관점",
        "8. 결과",
        "텍스트 조건의 효과는 frame prediction에서 먼저 확인된다.",
    )
    deck.image(slide, TEXT_DELTA_FIG, 0.75, 1.45, 6.2, 3.35)
    deck.card(slide, "일반적 경향", "텍스트 조건은 low-resolution 또는 rollout처럼 semantic cue가 사라지기 쉬운 환경에서 LPIPS/DreamSim/FID를 낮추는 경향을 보였습니다.", 7.35, 1.55, 4.7, 1.45, fill=LIGHT_GREEN, accent=GREEN)
    deck.card(slide, "주의", "frame prediction 개선이 planning ATE 개선으로 자동 연결되지는 않습니다. CEM score, action sampler, goal matching, 후반 trajectory outlier가 planning 결과를 바꿀 수 있습니다.", 7.35, 3.35, 4.7, 1.8, fill=LIGHT_ORANGE, accent=ORANGE)

    # 22
    slide = deck.slide(
        "Planning 결과 해석: 왜 metric이 엇갈리나",
        "8. 결과",
        "B224 Text에서 ATE와 RPE가 다르게 움직이는 현상.",
    )
    deck.bullets(
        slide,
        [
            "ATE는 전체 trajectory의 global drift에 민감합니다.",
            "RPE는 local motion 변화의 상대 일관성에 더 가깝습니다.",
            "B224 Text N120은 RPE가 N32보다 좋아졌지만, 후반 샘플의 큰 drift 때문에 ATE가 악화되었습니다.",
            "따라서 “text가 navigation을 무조건 개선한다”가 아니라 “어떤 metric/setting에서 개선되는지”로 읽어야 합니다.",
        ],
        0.85,
        1.65,
        6.2,
        2.6,
        size=14,
    )
    deck.metric_card(slide, "B224 Text ATE", "1.293", "N32 대비 +2.56%, paper 대비 +14.42%", 7.65, 1.6, 2.25, 1.45, LIGHT_RED, RED)
    deck.metric_card(slide, "B224 Text RPE", "0.350", "N32 대비 -1.17%, paper와 거의 동률", 10.15, 1.6, 2.25, 1.45, LIGHT_GREEN, GREEN)
    deck.card(slide, "실험 결론", "원 논문 CEM 설정으로 바꾸면 evaluation fairness는 좋아지지만, 모든 text-conditioned variant의 ATE가 자동으로 좋아지는 것은 아닙니다.", 7.65, 3.55, 4.75, 1.4, fill=WHITE, accent=BLUE)

    # 23
    slide = deck.slide(
        "End-to-End 실행 Workflow",
        "9. 실습",
        "처음부터 결과 JSON까지 가는 실전 순서.",
    )
    deck.flow(
        slide,
        [
            ("prepare", "dataset/splits\ntext cache"),
            ("smoke", "dataset load\n1-sample forward"),
            ("train", "CDiT checkpoint"),
            ("infer", "GT / prediction frames"),
            ("evaluate", "LPIPS/Dream/FID"),
            ("plan_eval", "ATE/RPE"),
        ],
        0.5,
        1.55,
        12.35,
        1.35,
        [WHITE, WHITE, WHITE, WHITE, WHITE, WHITE],
    )
    deck.code(
        slide,
        """docker build -t nwm:cu126 .
./scripts/docker/nwm-start.sh
./scripts/docker/nwm-run.sh "python tests/smoke/recon_smoke_test.py --skip-forward"
./scripts/docker/nwm-run.sh "python scripts/train.py --config configs/experiment/nwm_cdit_b_recon_raw_text_dense.yaml"
./scripts/docker/nwm-run.sh "python scripts/infer.py --exp configs/experiment/nwm_cdit_b_recon_raw_text_dense.yaml --datasets recon --eval_type time"
./scripts/docker/nwm-run.sh "python scripts/evaluate.py --datasets recon --gt_dir <gt> --exp_dir <pred>"
./scripts/docker/nwm-run.sh "python scripts/plan_eval.py --exp configs/experiment/nwm_cdit_b_recon_raw_text_dense.yaml --datasets recon --num_samples 120 --topk 5 --num_repeat_eval 3 --opt_steps 1" """,
        0.85,
        3.55,
        11.65,
        2.35,
        size=7,
    )

    # 24
    slide = deck.slide(
        "Smoke Test와 Debugging",
        "9. 실습",
        "큰 학습/평가 전에 확인해야 하는 작은 검증.",
    )
    deck.card(slide, "Dataset load", "`python tests/smoke/recon_smoke_test.py --skip-forward`\n\n이미지, trajectory metadata, action 계산, text cache shape를 먼저 확인합니다.", 0.85, 1.65, 3.75, 2.0, fill=LIGHT_BLUE, accent=BLUE)
    deck.card(slide, "One forward", "`python tests/smoke/recon_smoke_test.py --horizon-steps 8`\n\nVAE, CDiT forward, diffusion wrapper가 함께 동작하는지 확인합니다.", 4.85, 1.65, 3.75, 2.0, fill=LIGHT_TEAL, accent=TEAL)
    deck.card(slide, "Runtime checks", "`nvidia-smi`, log tail, result JSON 존재 여부, `verify_*` 스크립트로 산출물 무결성을 확인합니다.", 8.85, 1.65, 3.75, 2.0, fill=LIGHT_ORANGE, accent=ORANGE)
    deck.bullets(
        slide,
        [
            "OOM이면 batch size, num_samples, cem_eval_chunk_size를 먼저 봅니다.",
            "checkpoint shape mismatch는 image_size/model/text_dim 차이일 가능성이 큽니다.",
            "planning metric은 JSON이 생성된 뒤에만 최종값으로 봅니다. log의 rolling average는 중간값입니다.",
        ],
        1.0,
        4.45,
        11.2,
        1.25,
        size=13,
    )

    # 25
    slide = deck.slide(
        "Artifacts와 실험 기록",
        "9. 실습",
        "실험 결과가 어디에 쌓이고 무엇을 commit/공유할지.",
    )
    deck.card(slide, "weights/", "학습 checkpoint\n`weights/checkpoints/<run_name>/<step>.pth.tar`", 0.85, 1.55, 3.55, 1.2, fill=LIGHT_BLUE, accent=BLUE)
    deck.card(slide, "logs/", "학습/비동기 실행 로그\nresolved config와 runtime 기록", 4.85, 1.55, 3.55, 1.2, fill=LIGHT_TEAL, accent=TEAL)
    deck.card(slide, "artifacts/", "평가 frame, planning JSON, plot, summary, deck", 8.85, 1.55, 3.55, 1.2, fill=LIGHT_ORANGE, accent=ORANGE)
    deck.code(
        slide,
        """artifacts/
├── bulk/train/        training sample dumps
├── bulk/eval/         GT and prediction frames
├── bulk/planning/     planning prediction/metric JSONs
├── summaries/         compact metric CSV/MD/PNG
├── profiling/         runtime/profiling figures
└── decks/             presentation outputs""",
        1.0,
        3.45,
        5.25,
        2.2,
        size=9,
    )
    deck.card(slide, "실험 재현성", "config, checkpoint, dataset split, command, GPU/VRAM, result JSON path를 같이 기록해야 나중에 결과를 다시 해석할 수 있습니다.", 6.95, 3.45, 5.0, 1.7, fill=WHITE, accent=GREEN)

    # 26
    slide = deck.slide(
        "코드 읽기 로드맵",
        "10. 공부법",
        "수업 후 혼자 repo를 공부할 때 추천 순서.",
    )
    deck.table(
        slide,
        ["Week", "읽을 파일", "질문"],
        [
            ["1", "README.md, docs/project_structure.md", "전체 pipeline을 말로 설명할 수 있나?"],
            ["2", "src/data/datasets/*.py", "sample이 어떤 tensor로 변환되나?"],
            ["3", "src/models/backbones/cdit.py", "condition vector가 block에 어떻게 들어가나?"],
            ["4", "scripts/train.py", "loss와 checkpoint가 어떻게 만들어지나?"],
            ["5", "scripts/infer.py, rollout.py", "time/rollout prediction 차이는?"],
            ["6", "cem_planner.py", "CEM이 action 후보를 어떻게 고르나?"],
        ],
        0.85,
        1.55,
        11.65,
        4.1,
        font_size=9,
    )
    deck.text(slide, "읽을 때 항상 “입력 tensor shape → condition → output artifact”를 추적하세요.", 0.95, 6.05, 11.3, 0.35, 14, True, DARK, align=PP_ALIGN.CENTER)

    # 27
    slide = deck.slide(
        "Lab 1: 데이터셋 한 샘플 해부",
        "10. 공부법",
        "첫 번째 실습 과제.",
    )
    deck.bullets(
        slide,
        [
            "`TrainingDataset.__getitem__`에 breakpoint 또는 print를 넣고 한 샘플을 꺼냅니다.",
            "`obs_image.shape`, `goal_pos.shape`, `rel_time`, `text_emb.shape`를 기록합니다.",
            "`_compute_actions`에서 global pose가 local x/y/yaw로 바뀌는 과정을 따라갑니다.",
            "text cache가 없는 trajectory에서 zero embedding이 들어가는지 확인합니다.",
        ],
        0.9,
        1.65,
        6.15,
        2.6,
        size=14,
    )
    deck.code(
        slide,
        """python - <<'PY'
from src.config import load_runtime_config
from src.data.transforms.image import build_transform
# config를 읽고 TrainingDataset 하나를 직접 생성해 shape를 출력해보는 연습
PY""",
        0.95,
        4.65,
        6.0,
        0.95,
        size=9,
    )
    deck.card(slide, "제출물", "샘플 tensor shape 표\n한 trajectory의 curr_time/goal_time 예시\n텍스트 embedding source 설명", 7.65, 1.8, 3.95, 2.0, fill=LIGHT_BLUE, accent=BLUE)

    # 28
    slide = deck.slide(
        "Lab 2: 모델 forward 따라가기",
        "10. 공부법",
        "두 번째 실습 과제.",
    )
    deck.bullets(
        slide,
        [
            "`CDiT.forward(x, t, y, x_cond, rel_t, text_emb)`의 각 인자 shape를 적습니다.",
            "`x_embedder`, `pos_embed`, `y_embedder`, `time_embedder`, `text_proj`를 차례로 확인합니다.",
            "`CDiTBlock.forward`에서 self-attn, cross-attn, MLP가 어떤 순서로 적용되는지 도식화합니다.",
            "text_emb=None일 때와 text_emb가 있을 때 condition vector c가 어떻게 달라지는지 비교합니다.",
        ],
        0.9,
        1.65,
        6.4,
        2.8,
        size=14,
    )
    deck.card(slide, "핵심 질문", "텍스트는 token attention으로 들어가는가, 아니면 global condition으로 들어가는가?", 7.8, 1.75, 3.9, 1.15, fill=LIGHT_ORANGE, accent=ORANGE)
    deck.card(slide, "정답", "현재 구현에서는 `text_proj(text_emb)`가 `c = t + time + action`에 더해지고, 이 c가 adaLN modulation에 사용됩니다.", 7.8, 3.2, 3.9, 1.55, fill=LIGHT_GREEN, accent=GREEN)

    # 29
    slide = deck.slide(
        "Lab 3: 평가 결과 재현",
        "10. 공부법",
        "세 번째 실습 과제.",
    )
    deck.bullets(
        slide,
        [
            "GT frame 생성과 prediction frame 생성을 분리해서 실행합니다.",
            "LPIPS/DreamSim/FID JSON을 만들고, horizon별 metric이 어떻게 변하는지 plot합니다.",
            "planning은 작은 `--max_eval_samples`로 먼저 smoke 후 full run으로 확장합니다.",
            "CEM 설정을 바꿔도 모델 checkpoint는 바뀌지 않는다는 점을 실험 로그로 확인합니다.",
        ],
        0.9,
        1.55,
        6.25,
        2.7,
        size=14,
    )
    deck.code(
        slide,
        """python scripts/plan_eval.py \\
  --exp configs/experiment/nwm_cdit_b_recon_raw_text_dense.yaml \\
  --datasets recon --ckp 0030000 \\
  --num_samples 120 --topk 5 --num_repeat_eval 3 \\
  --opt_steps 1 --action_sampler repeat --rollout_stride 1""",
        7.45,
        1.75,
        4.75,
        1.7,
        size=8,
    )
    deck.card(slide, "제출물", "명령어, checkpoint, output JSON path, ATE/RPE/Pos/Yaw, GPU와 runtime 기록", 7.45, 4.05, 4.75, 1.2, fill=WHITE, accent=BLUE)

    # 30
    slide = deck.slide(
        "시험에 나올 만한 질문",
        "10. 공부법",
        "프로젝트 이해도를 스스로 점검하는 질문.",
    )
    deck.bullets(
        slide,
        [
            "NWM은 policy인가, world model인가? 왜 그렇게 말할 수 있는가?",
            "CDiT에서 self-attention과 cross-attention은 각각 무엇을 보는가?",
            "text conditioning이 모델 architecture에 들어가는 정확한 위치는 어디인가?",
            "time prediction, rollout, planning evaluation은 무엇이 다르고 왜 모두 필요한가?",
            "LPIPS/DreamSim/FID와 ATE/RPE는 서로 어떤 종류의 성능을 측정하는가?",
            "CEM 설정 변경은 왜 재학습이 아니라 평가 방법 변경인가?",
            "B224 Text paper-style CEM 결과에서 ATE와 RPE 결론이 왜 다르게 나왔는가?",
        ],
        0.9,
        1.55,
        11.3,
        4.6,
        size=14,
    )

    # 31
    slide = deck.slide(
        "핵심 Takeaways",
        "11. 정리",
        "프로젝트를 한 장으로 요약.",
    )
    deck.card(slide, "1. 문제", "현재 관측과 행동으로 미래 시각 상태를 예측해 navigation planning을 돕는 world model입니다.", 0.8, 1.55, 3.55, 1.35, fill=LIGHT_BLUE, accent=BLUE)
    deck.card(slide, "2. 방법", "VAE latent diffusion + CDiT + action/time/text conditioning으로 future frame latent를 denoise합니다.", 4.85, 1.55, 3.55, 1.35, fill=LIGHT_TEAL, accent=TEAL)
    deck.card(slide, "3. 평가", "이미지 품질은 LPIPS/DreamSim/FID, navigation은 CEM planning의 ATE/RPE로 봅니다.", 8.9, 1.55, 3.55, 1.35, fill=LIGHT_ORANGE, accent=ORANGE)
    deck.card(slide, "4. 현재 결과", "Paper-style CEM full 3/3 완료. B224 No-Text는 paper NWM과 거의 동률, B128 Text는 N32 대비 개선, B224 Text는 ATE 악화/RPE 동률권입니다.", 0.8, 3.45, 5.65, 1.55, fill=WHITE, accent=GREEN)
    deck.card(slide, "5. 다음 연구 질문", "텍스트가 planning ATE에 손해를 주는 후반 trajectory outlier를 줄이려면 cost function, text source, ranker, CEM sampler를 어떻게 바꿔야 하는가?", 6.85, 3.45, 5.65, 1.55, fill=WHITE, accent=PURPLE)
    deck.text(slide, "이제 코드를 볼 때는 항상 “data → latent → condition → denoise → decode → metric” 흐름으로 추적하면 됩니다.", 0.9, 6.05, 11.5, 0.45, 15, True, DARK, align=PP_ALIGN.CENTER)

    deck.prs.save(PPTX_PATH)
    write_notes(deck.notes)
    maybe_export_pdf()


def write_notes(notes: list[tuple[str, str]]) -> None:
    lines = [
        "# NWM Project Course Deck Notes",
        "",
        "이 노트는 PPT 각 슬라이드의 강의 의도를 짧게 정리한 보조 자료입니다.",
        "",
    ]
    for idx, (title, note) in enumerate(notes, start=1):
        lines.append(f"## {idx:02d}. {title}")
        lines.append("")
        lines.append(textwrap.fill(note or "슬라이드 내용을 기준으로 설명합니다.", width=96))
        lines.append("")
    NOTES_PATH.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_deck()
    print(f"Wrote {PPTX_PATH.relative_to(ROOT)}")
    if PDF_PATH.exists():
        print(f"Wrote {PDF_PATH.relative_to(ROOT)}")
    print(f"Wrote {NOTES_PATH.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
