#!/usr/bin/env python3
"""Build the post-0427 NWM text-conditioned results deck.

The script collects local experiment summaries, adds reference metrics from the
Navigation World Models paper, renders comparison figures, and writes an
editable PowerPoint deck plus a compact notes file.
"""

from __future__ import annotations

from pathlib import Path
import shutil
import subprocess
import textwrap

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "artifacts" / "decks" / "nwm_text_conditioned_update"
FIG_DIR = OUT_DIR / "figures"
DATA_DIR = OUT_DIR / "data"
PPTX_PATH = OUT_DIR / "nwm_text_conditioned_update.pptx"
NOTES_PATH = OUT_DIR / "nwm_text_conditioned_update_notes.md"
PAPER_CEM_SUMMARY_DIR = ROOT / "artifacts" / "summaries" / "planning" / "paper_cem_b_selected_n120_rep3"
PAPER_CEM_CSV = PAPER_CEM_SUMMARY_DIR / "metrics.csv"
PAPER_CEM_FIG = PAPER_CEM_SUMMARY_DIR / "ate_rpe_comparison.png"

S_PLANNING_CSV = (
    ROOT / "artifacts" / "summaries" / "planning" / "text_vs_notext_128_224_n32" / "metrics.csv"
)
B_PLANNING_CSV = (
    ROOT / "artifacts" / "summaries" / "planning" / "text_vs_notext_128_224_b_n32" / "metrics.csv"
)
S_PLANNING_BARS = (
    ROOT / "artifacts" / "summaries" / "planning" / "text_vs_notext_128_224_n32" / "metrics_bars.png"
)
B_PLANNING_BARS = (
    ROOT / "artifacts" / "summaries" / "planning" / "text_vs_notext_128_224_b_n32" / "metrics_bars.png"
)
S_TIME_CURVE = ROOT / "artifacts" / "profiling" / "recon_time_over_checkpoints" / "01_recon_time_over_checkpoints.png"
B_TIME_CURVE = (
    ROOT
    / "artifacts"
    / "profiling"
    / "recon_time_over_checkpoints_cdit_b"
    / "01_recon_time_over_checkpoints.png"
)

FONT = "Noto Sans CJK KR"
TITLE = RGBColor(22, 28, 36)
BODY = RGBColor(52, 60, 70)
MUTED = RGBColor(103, 112, 124)
BG = RGBColor(247, 248, 250)
NAVY = RGBColor(34, 74, 116)
TEAL = RGBColor(21, 128, 121)
ORANGE = RGBColor(199, 108, 32)
RED = RGBColor(178, 60, 54)
GREEN = RGBColor(51, 135, 78)
LIGHT_BLUE = RGBColor(224, 235, 246)
LIGHT_TEAL = RGBColor(222, 242, 239)
LIGHT_ORANGE = RGBColor(250, 235, 220)
LIGHT_GRAY = RGBColor(234, 237, 241)
WHITE = RGBColor(255, 255, 255)


def ensure_dirs() -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    DATA_DIR.mkdir(parents=True, exist_ok=True)


def frame_prediction_data() -> pd.DataFrame:
    rows = [
        {
            "model": "CDiT-S",
            "resolution": 224,
            "condition": "No-Text",
            "lpips4": 0.36985841393470764,
            "dream4": 0.16149090230464935,
            "fid4": 33.993408203125,
            "lpips16": 0.46318289637565613,
            "dream16": 0.2286277711391449,
            "fid16": 40.388671875,
        },
        {
            "model": "CDiT-S",
            "resolution": 128,
            "condition": "No-Text",
            "lpips4": 0.30775272846221924,
            "dream4": 0.19891947507858276,
            "fid4": 33.02520751953125,
            "lpips16": 0.40308988094329834,
            "dream16": 0.2632315456867218,
            "fid16": 38.088470458984375,
        },
        {
            "model": "CDiT-S",
            "resolution": 128,
            "condition": "Text",
            "lpips4": 0.29237666726112366,
            "dream4": 0.18870626389980316,
            "fid4": 33.41960144042969,
            "lpips16": 0.38733357191085815,
            "dream16": 0.24954962730407715,
            "fid16": 36.69294738769531,
        },
        {
            "model": "CDiT-S",
            "resolution": 224,
            "condition": "Text",
            "lpips4": 0.3657650947570801,
            "dream4": 0.15901166200637817,
            "fid4": 33.07463073730469,
            "lpips16": 0.4576801061630249,
            "dream16": 0.22435526549816132,
            "fid16": 38.1861572265625,
        },
        {
            "model": "CDiT-B",
            "resolution": 224,
            "condition": "No-Text",
            "lpips4": 0.3238004744052887,
            "dream4": 0.12083755433559418,
            "fid4": 27.606689453125,
            "lpips16": 0.43125712871551514,
            "dream16": 0.1867092102766037,
            "fid16": 32.76750183105469,
        },
        {
            "model": "CDiT-B",
            "resolution": 128,
            "condition": "No-Text",
            "lpips4": 0.2701418101787567,
            "dream4": 0.1632404327392578,
            "fid4": 29.580108642578125,
            "lpips16": 0.37984219193458557,
            "dream16": 0.22864902019500732,
            "fid16": 33.6153564453125,
        },
        {
            "model": "CDiT-B",
            "resolution": 128,
            "condition": "Text",
            "lpips4": 0.2619237005710602,
            "dream4": 0.15621471405029297,
            "fid4": 28.942886352539062,
            "lpips16": 0.36699816584587097,
            "dream16": 0.21926714479923248,
            "fid16": 33.78776550292969,
        },
        {
            "model": "CDiT-B",
            "resolution": 224,
            "condition": "Text",
            "lpips4": 0.3381218910217285,
            "dream4": 0.12444277852773666,
            "fid4": 28.436111450195312,
            "lpips16": 0.43855586647987366,
            "dream16": 0.18868020176887512,
            "fid16": 31.9881591796875,
        },
    ]
    df = pd.DataFrame(rows)
    df["variant"] = df.apply(
        lambda r: f"{r['model']} {int(r['resolution'])} {r['condition']}", axis=1
    )
    return df


def load_planning_data() -> pd.DataFrame:
    s = pd.read_csv(S_PLANNING_CSV).rename(
        columns={
            "resolution": "resolution",
            "recon_ate": "ate",
            "recon_rpe_trans": "rpe_trans",
            "recon_pos_diff_norm": "pos_error",
            "recon_yaw_diff_norm": "yaw_error",
            "total_time": "time_s",
        }
    )
    s["model"] = "CDiT-S"
    s = s[["model", "resolution", "condition", "ate", "rpe_trans", "pos_error", "yaw_error", "time_s"]]

    b = pd.read_csv(B_PLANNING_CSV).rename(columns={"image": "resolution"})
    b["model"] = "CDiT-B"
    b = b[["model", "resolution", "condition", "ate", "rpe_trans", "pos_error", "yaw_error", "time_s"]]

    df = pd.concat([s, b], ignore_index=True)
    df["resolution"] = df["resolution"].astype(int)
    df["variant"] = df.apply(
        lambda r: f"{r['model']} {int(r['resolution'])} {r['condition']}", axis=1
    )
    return df.sort_values(["model", "resolution", "condition"]).reset_index(drop=True)


def paper_planning_data() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"method": "Paper Forward", "ate": 1.92, "rpe_trans": 0.54, "source": "Table 7"},
            {"method": "Paper GNM", "ate": 1.87, "rpe_trans": 0.73, "source": "Table 7"},
            {"method": "Paper NoMaD", "ate": 1.95, "rpe_trans": 0.53, "source": "Table 7"},
            {
                "method": "Paper NWM+NoMaD x16",
                "ate": 1.88,
                "rpe_trans": 0.51,
                "source": "Table 7",
            },
            {
                "method": "Paper NWM+NoMaD x32",
                "ate": 1.79,
                "rpe_trans": 0.49,
                "source": "Table 7",
            },
            {"method": "Paper NWM only", "ate": 1.13, "rpe_trans": 0.35, "source": "Table 7"},
        ]
    )


def paper_prediction_data() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "method": "Paper time only",
                "lpips4": 0.760,
                "dream4": 0.783,
                "psnr4": 7.839,
                "source": "Table 1",
            },
            {
                "method": "Paper action only",
                "lpips4": 0.318,
                "dream4": 0.100,
                "psnr4": 14.858,
                "source": "Table 1",
            },
            {
                "method": "Paper action+time",
                "lpips4": 0.295,
                "dream4": 0.091,
                "psnr4": 15.343,
                "source": "Table 1",
            },
            {
                "method": "Paper 4-goal/context",
                "lpips4": 0.296,
                "dream4": 0.091,
                "psnr4": 15.331,
                "source": "Table 1",
            },
        ]
    )


def text_delta(df: pd.DataFrame, metrics: list[str]) -> pd.DataFrame:
    rows = []
    for (model, resolution), group in df.groupby(["model", "resolution"]):
        if set(group["condition"]) != {"No-Text", "Text"}:
            continue
        base = group[group["condition"] == "No-Text"].iloc[0]
        text = group[group["condition"] == "Text"].iloc[0]
        row = {"model": model, "resolution": resolution, "label": f"{model} {resolution}"}
        for metric in metrics:
            row[metric] = (text[metric] - base[metric]) / base[metric] * 100.0
        rows.append(row)
    return pd.DataFrame(rows).sort_values(["model", "resolution"]).reset_index(drop=True)


def save_data() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    frame = frame_prediction_data()
    planning = load_planning_data()
    paper_plan = paper_planning_data()
    paper_pred = paper_prediction_data()

    frame.to_csv(DATA_DIR / "frame_prediction_metrics.csv", index=False)
    planning.to_csv(DATA_DIR / "planning_metrics.csv", index=False)
    paper_plan.to_csv(DATA_DIR / "paper_planning_metrics.csv", index=False)
    paper_pred.to_csv(DATA_DIR / "paper_prediction_metrics.csv", index=False)
    text_delta(frame, ["lpips4", "dream4", "fid4", "lpips16", "dream16", "fid16"]).to_csv(
        DATA_DIR / "frame_prediction_text_delta_percent.csv", index=False
    )
    text_delta(planning, ["ate", "rpe_trans", "pos_error", "yaw_error"]).to_csv(
        DATA_DIR / "planning_text_delta_percent.csv", index=False
    )
    return frame, planning, paper_plan, paper_pred


def set_plot_style() -> None:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "axes.titlesize": 13,
            "axes.labelsize": 10,
            "xtick.labelsize": 8,
            "ytick.labelsize": 9,
            "figure.dpi": 140,
            "savefig.dpi": 180,
        }
    )


def colors_for_variants(labels: list[str]) -> list[str]:
    colors = []
    for label in labels:
        if label.startswith("Paper"):
            colors.append("#5b6470")
        elif "CDiT-S" in label and "Text" in label:
            colors.append("#1f9e89")
        elif "CDiT-S" in label:
            colors.append("#74c7b8")
        elif "CDiT-B" in label and "Text" in label:
            colors.append("#d98732")
        else:
            colors.append("#95a6bd")
    return colors


def add_value_labels(ax, bars, fmt: str = "{:.3f}", pad: float = 0.006) -> None:
    for bar in bars:
        height = bar.get_height()
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            height + pad,
            fmt.format(height),
            ha="center",
            va="bottom",
            fontsize=7,
            color="#303844",
            rotation=90 if len(bars) > 8 else 0,
        )


def render_bar_metric(
    ax,
    labels: list[str],
    values: list[float],
    title: str,
    ylabel: str,
    paper_value: float | None = None,
) -> None:
    bars = ax.bar(range(len(labels)), values, color=colors_for_variants(labels), width=0.72)
    ax.set_title(title, loc="left", fontweight="bold")
    ax.set_ylabel(ylabel)
    ax.set_xticks(range(len(labels)))
    ax.set_xticklabels(labels, rotation=45, ha="right")
    ax.grid(axis="y", color="#d9dde3", linewidth=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    if paper_value is not None:
        ax.axhline(paper_value, color="#b33d36", linestyle="--", linewidth=1.2)
        ax.text(
            len(labels) - 0.2,
            paper_value,
            f"paper {paper_value:.3f}",
            ha="right",
            va="bottom",
            fontsize=8,
            color="#8d2f2a",
        )
    add_value_labels(ax, bars)


def make_figures(frame: pd.DataFrame, planning: pd.DataFrame, paper_plan: pd.DataFrame) -> dict[str, Path]:
    set_plot_style()
    outputs: dict[str, Path] = {}

    pred_labels = ["Paper action+time"] + frame["variant"].tolist()
    lpips_values = [0.295] + frame["lpips4"].tolist()
    dream_values = [0.091] + frame["dream4"].tolist()
    fig, axs = plt.subplots(1, 2, figsize=(12.8, 5.6))
    render_bar_metric(axs[0], pred_labels, lpips_values, "4s LPIPS", "lower is better", 0.295)
    render_bar_metric(axs[1], pred_labels, dream_values, "4s DreamSim", "lower is better", 0.091)
    fig.suptitle("Frame prediction vs. NWM paper reference", x=0.02, ha="left", fontsize=16, fontweight="bold")
    fig.tight_layout(rect=[0, 0.02, 1, 0.93])
    outputs["paper_prediction"] = FIG_DIR / "paper_prediction_comparison.png"
    fig.savefig(outputs["paper_prediction"], bbox_inches="tight")
    plt.close(fig)

    plan_labels = ["Paper NWM only"] + planning["variant"].tolist()
    ate_values = [1.13] + planning["ate"].tolist()
    rpe_values = [0.35] + planning["rpe_trans"].tolist()
    fig, axs = plt.subplots(1, 2, figsize=(12.8, 5.6))
    render_bar_metric(axs[0], plan_labels, ate_values, "Planning ATE", "lower is better", 1.13)
    render_bar_metric(axs[1], plan_labels, rpe_values, "Planning RPE trans", "lower is better", 0.35)
    fig.suptitle("Planning metrics vs. NWM paper reference", x=0.02, ha="left", fontsize=16, fontweight="bold")
    fig.tight_layout(rect=[0, 0.02, 1, 0.93])
    outputs["paper_planning"] = FIG_DIR / "paper_planning_comparison.png"
    fig.savefig(outputs["paper_planning"], bbox_inches="tight")
    plt.close(fig)

    pred_delta = text_delta(frame, ["lpips4", "dream4", "fid4", "lpips16", "dream16", "fid16"])
    fig, ax = plt.subplots(figsize=(10.8, 4.8))
    data = pred_delta[["lpips4", "dream4", "fid4", "lpips16", "dream16", "fid16"]].to_numpy()
    im = ax.imshow(data, cmap="RdYlGn_r", vmin=-15, vmax=15)
    ax.set_yticks(range(len(pred_delta)))
    ax.set_yticklabels(pred_delta["label"])
    ax.set_xticks(range(data.shape[1]))
    ax.set_xticklabels(["LPIPS 4s", "Dream 4s", "FID 4s", "LPIPS 16s", "Dream 16s", "FID 16s"], rotation=25, ha="right")
    for i in range(data.shape[0]):
        for j in range(data.shape[1]):
            ax.text(j, i, f"{data[i, j]:+.1f}%", ha="center", va="center", fontsize=9, color="#1d252f")
    ax.set_title("Text minus No-Text: frame prediction delta", loc="left", fontweight="bold")
    ax.text(
        0,
        -0.75,
        "Negative is better because all metrics are lower-is-better.",
        ha="left",
        va="bottom",
        fontsize=9,
        color="#596270",
    )
    fig.colorbar(im, ax=ax, shrink=0.78, label="% delta")
    fig.tight_layout()
    outputs["prediction_delta"] = FIG_DIR / "prediction_text_delta_heatmap.png"
    fig.savefig(outputs["prediction_delta"], bbox_inches="tight")
    plt.close(fig)

    plan_delta = text_delta(planning, ["ate", "rpe_trans", "pos_error", "yaw_error"])
    fig, ax = plt.subplots(figsize=(9.4, 4.6))
    data = plan_delta[["ate", "rpe_trans", "pos_error", "yaw_error"]].to_numpy()
    im = ax.imshow(data, cmap="RdYlGn_r", vmin=-25, vmax=25)
    ax.set_yticks(range(len(plan_delta)))
    ax.set_yticklabels(plan_delta["label"])
    ax.set_xticks(range(data.shape[1]))
    ax.set_xticklabels(["ATE", "RPE", "Position", "Yaw"], rotation=0)
    for i in range(data.shape[0]):
        for j in range(data.shape[1]):
            ax.text(j, i, f"{data[i, j]:+.1f}%", ha="center", va="center", fontsize=10, color="#1d252f")
    ax.set_title("Text minus No-Text: planning delta", loc="left", fontweight="bold")
    ax.text(
        0,
        -0.7,
        "Negative is better. B shows more consistent RPE/Yaw gains; ATE does not reliably improve.",
        ha="left",
        va="bottom",
        fontsize=9,
        color="#596270",
    )
    fig.colorbar(im, ax=ax, shrink=0.78, label="% delta")
    fig.tight_layout()
    outputs["planning_delta"] = FIG_DIR / "planning_text_delta_heatmap.png"
    fig.savefig(outputs["planning_delta"], bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(10.2, 4.8))
    runtime = planning.copy()
    runtime["hours"] = runtime["time_s"] / 3600.0
    labels = runtime["variant"].tolist()
    bars = ax.bar(range(len(labels)), runtime["hours"], color=colors_for_variants(labels), width=0.72)
    ax.set_title("Full 100-sample planning eval runtime", loc="left", fontweight="bold")
    ax.set_ylabel("hours")
    ax.set_xticks(range(len(labels)))
    ax.set_xticklabels(labels, rotation=35, ha="right")
    ax.grid(axis="y", color="#d9dde3", linewidth=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    for bar in bars:
        h = bar.get_height()
        ax.text(bar.get_x() + bar.get_width() / 2, h + 0.03, f"{h:.2f}h", ha="center", va="bottom", fontsize=8)
    fig.tight_layout()
    outputs["runtime"] = FIG_DIR / "planning_runtime.png"
    fig.savefig(outputs["runtime"], bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(9.0, 5.0))
    x = np.arange(4)
    order = [
        ("CDiT-S", 128),
        ("CDiT-B", 128),
        ("CDiT-S", 224),
        ("CDiT-B", 224),
    ]
    no_text = []
    text = []
    labels = []
    for model, resolution in order:
        group = planning[(planning["model"] == model) & (planning["resolution"] == resolution)]
        no_text.append(float(group[group["condition"] == "No-Text"]["ate"].iloc[0]))
        text.append(float(group[group["condition"] == "Text"]["ate"].iloc[0]))
        labels.append(f"{model}\n{resolution}")
    ax.bar(x - 0.18, no_text, width=0.36, label="No-Text", color="#95a6bd")
    ax.bar(x + 0.18, text, width=0.36, label="Text", color="#d98732")
    ax.axhline(1.13, linestyle="--", color="#b33d36", linewidth=1.1, label="Paper NWM only")
    ax.set_xticks(x)
    ax.set_xticklabels(labels)
    ax.set_ylabel("ATE")
    ax.set_title("S/B planning ATE by resolution", loc="left", fontweight="bold")
    ax.grid(axis="y", color="#d9dde3", linewidth=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(ncol=3, fontsize=9, loc="upper right")
    fig.tight_layout()
    outputs["planning_ate_grouped"] = FIG_DIR / "planning_ate_grouped.png"
    fig.savefig(outputs["planning_ate_grouped"], bbox_inches="tight")
    plt.close(fig)

    return outputs


def blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    return slide


def add_textbox(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    size: int = 20,
    color: RGBColor = BODY,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_header(slide, title: str, kicker: str | None = None) -> None:
    if kicker:
        add_textbox(slide, 0.6, 0.28, 11.8, 0.28, kicker, size=10, color=MUTED, bold=True)
    add_textbox(slide, 0.6, 0.52, 11.8, 0.58, title, size=26, color=TITLE, bold=True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.16), Inches(12.1), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = LIGHT_GRAY
    line.line.fill.background()


def add_bullets(slide, left: float, top: float, width: float, height: float, bullets: list[str], size: int = 17) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.word_wrap = True
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = BODY
        p.space_after = Pt(7)


def add_source(slide, text: str) -> None:
    add_textbox(slide, 0.62, 7.12, 12.0, 0.22, text, size=8, color=MUTED)


def add_round_rect(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    fill: RGBColor = WHITE,
    line: RGBColor = LIGHT_GRAY,
    radius_shape=MSO_SHAPE.ROUNDED_RECTANGLE,
):
    shape = slide.shapes.add_shape(radius_shape, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.8)
    return shape


def add_stat_card(slide, left: float, top: float, width: float, label: str, value: str, note: str, color: RGBColor) -> None:
    add_round_rect(slide, left, top, width, 1.18, fill=WHITE, line=LIGHT_GRAY)
    add_textbox(slide, left + 0.18, top + 0.14, width - 0.36, 0.22, label, size=9, color=MUTED, bold=True)
    add_textbox(slide, left + 0.18, top + 0.39, width - 0.36, 0.34, value, size=22, color=color, bold=True)
    add_textbox(slide, left + 0.18, top + 0.78, width - 0.36, 0.27, note, size=10, color=BODY)


def add_table(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    rows: list[list[str]],
    header_fill: RGBColor = LIGHT_BLUE,
    font_size: int = 10,
) -> None:
    table_shape = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(left), Inches(top), Inches(width), Inches(height)
    )
    table = table_shape.table
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.margin_left = Inches(0.05)
            cell.margin_right = Inches(0.05)
            fill = cell.fill
            fill.solid()
            fill.fore_color.rgb = header_fill if r_idx == 0 else WHITE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER if c_idx > 1 or r_idx == 0 else PP_ALIGN.LEFT
                for run in paragraph.runs:
                    run.font.name = FONT
                    run.font.size = Pt(font_size if r_idx else font_size + 1)
                    run.font.bold = r_idx == 0
                    run.font.color.rgb = TITLE if r_idx == 0 else BODY


def add_picture(slide, path: Path, left: float, top: float, width: float | None = None, height: float | None = None) -> None:
    if width is None and height is None:
        slide.shapes.add_picture(str(path), Inches(left), Inches(top))
    elif width is None:
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), height=Inches(height))
    elif height is None:
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))
    else:
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width), height=Inches(height))


def add_metric_table(slide, df: pd.DataFrame, metrics: list[str], left: float, top: float, width: float, height: float) -> None:
    headers = ["Variant"] + metrics
    rows = [headers]
    for _, row in df.iterrows():
        values = [row["variant"]]
        for metric in metrics:
            values.append(f"{row[metric]:.3f}")
        rows.append(values)
    add_table(slide, left, top, width, height, rows, font_size=8)


def table_from_planning(planning: pd.DataFrame, model: str) -> list[list[str]]:
    rows = [["Res", "Cond", "ATE", "RPE", "Pos", "Yaw", "Time"]]
    subset = planning[planning["model"] == model].sort_values(["resolution", "condition"])
    for _, row in subset.iterrows():
        rows.append(
            [
                str(int(row["resolution"])),
                row["condition"],
                f"{row['ate']:.3f}",
                f"{row['rpe_trans']:.3f}",
                f"{row['pos_error']:.3f}",
                f"{row['yaw_error']:.3f}",
                f"{row['time_s'] / 3600:.2f}h",
            ]
        )
    return rows


def paper_cem_status() -> tuple[pd.DataFrame | None, list[str]]:
    if not PAPER_CEM_CSV.exists():
        return None, []
    df = pd.read_csv(PAPER_CEM_CSV)
    full = df[(df["setting"] == "N120/K5/rep3/OPT1") & (df["samples"].astype(str) == "100")]
    done_full = full[full["status"] == "done"].copy()
    bullets = [
        "N120/rep3는 모델 재학습이 아니라 CEM planning search budget 변경이다.",
        f"Full N120/rep3 완료: {len(done_full)}/{len(full)}.",
    ]
    if not done_full.empty:
        best_ate = done_full.loc[done_full["ate"].astype(float).idxmin()]
        best_rpe = done_full.loc[done_full["rpe_trans"].astype(float).idxmin()]
        bullets.extend(
            [
                f"Best N120 ATE: {best_ate['label']} = {float(best_ate['ate']):.3f} ({float(best_ate['ate_delta_vs_paper_pct']):+.1f}% vs paper).",
                f"Best N120 RPE: {best_rpe['label']} = {float(best_rpe['rpe_trans']):.3f} ({float(best_rpe['rpe_delta_vs_paper_pct']):+.1f}% vs paper).",
            ]
        )
    else:
        smoke = df[df["setting"] == "N120/K5/rep3/OPT1 smoke"]
        if not smoke.empty:
            bullets.append("현재는 smoke 결과와 기존 N32 baseline만 포함되어 있다.")
    return df, bullets


def build_deck(frame: pd.DataFrame, planning: pd.DataFrame, paper_plan: pd.DataFrame, figures: dict[str, Path]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1. Title
    slide = blank_slide(prs)
    add_textbox(slide, 0.75, 0.72, 11.8, 0.44, "Navigation World Models", size=18, color=NAVY, bold=True)
    add_textbox(
        slide,
        0.75,
        1.35,
        11.6,
        1.25,
        "Text-Conditioned NWM\nPost-0427 Training & Evaluation",
        size=36,
        color=TITLE,
        bold=True,
    )
    add_textbox(
        slide,
        0.78,
        2.86,
        10.5,
        0.58,
        "CDiT-S/B, 128/224, text-dense variants, frame prediction and CEM planning results",
        size=18,
        color=BODY,
    )
    add_stat_card(slide, 0.8, 4.35, 2.45, "Best local ATE", "1.162", "CDiT-B 224 No-Text", NAVY)
    add_stat_card(slide, 3.45, 4.35, 2.45, "Best local RPE", "0.354", "CDiT-B 224 Text", TEAL)
    add_stat_card(slide, 6.1, 4.35, 2.45, "Paper NWM", "1.13 / 0.35", "ATE / RPE reference", ORANGE)
    add_stat_card(slide, 8.75, 4.35, 2.45, "Main finding", "Text ≠ ATE", "semantic gain, geometry gap", RED)
    add_source(slide, "Sources: Navigation World Models.pdf; 0427 NWM Text-Conditioned deck; local artifacts under artifacts/profiling and artifacts/summaries/planning.")

    # 2. After 0427
    slide = blank_slide(prs)
    add_header(slide, "0427 이후 무엇을 했나", "Scope")
    add_bullets(
        slide,
        0.75,
        1.45,
        6.15,
        4.9,
        [
            "0427 자료의 핵심 질문: 저해상도 visual world model에서 텍스트가 semantic prior로 작동하는가.",
            "이후 CDiT-S뿐 아니라 CDiT-B까지 확장해 128/224, No-Text/Text-Dense 조합을 평가했다.",
            "평가는 future-frame prediction과 navigation planning으로 나눴다.",
            "원 논문 NWM의 RECON 지표를 기준선으로 넣어, 논문 대비 현재 위치를 수치로 비교했다.",
        ],
        size=18,
    )
    timeline = [
        ("0427", "idea + early\nprediction eval", LIGHT_BLUE),
        ("After", "S/B training\n128/224 variants", LIGHT_TEAL),
        ("Now", "frame + planning\npaper comparison", LIGHT_ORANGE),
    ]
    for idx, (label, desc, fill) in enumerate(timeline):
        left = 7.35 + idx * 1.72
        add_round_rect(slide, left, 2.0, 1.36, 1.55, fill=fill)
        add_textbox(slide, left + 0.15, 2.18, 1.05, 0.26, label, size=14, color=TITLE, bold=True, align=PP_ALIGN.CENTER)
        add_textbox(slide, left + 0.12, 2.63, 1.12, 0.58, desc, size=11, color=BODY, align=PP_ALIGN.CENTER)
    add_textbox(slide, 7.42, 4.32, 4.85, 0.9, "결론 축", size=18, color=TITLE, bold=True)
    add_bullets(
        slide,
        7.42,
        4.78,
        4.85,
        1.7,
        [
            "Text는 frame metric에서 자주 이득을 보인다.",
            "Planning ATE는 그 이득을 그대로 따라오지 않는다.",
            "CDiT-B는 S보다 yaw/RPE 개선이 더 일관적이다.",
        ],
        size=15,
    )
    add_source(slide, "Source: 0427 deck slides 17-20; local S/B summary files.")

    # 3. Paper recap
    slide = blank_slide(prs)
    add_header(slide, "원 논문의 NWM 기준선", "Paper Reference")
    add_bullets(
        slide,
        0.72,
        1.5,
        5.6,
        4.85,
        [
            "NWM은 과거 관측과 navigation action을 조건으로 미래 관측을 생성하는 controllable video generation model이다.",
            "CDiT는 context frame 수에 대해 선형 복잡도를 갖도록 설계되었다.",
            "논문은 image metric, video distribution metric, navigation metric을 함께 보고한다.",
            "우리 비교에서는 현재 재현 가능한 LPIPS/DreamSim/FID와 ATE/RPE에 초점을 맞춘다.",
        ],
        size=17,
    )
    paper_rows = [
        ["Metric", "Paper best/reference"],
        ["4s LPIPS", "0.295 action+time"],
        ["4s DreamSim", "0.091 action+time"],
        ["4s PSNR", "15.343 action+time"],
        ["Planning ATE", "1.13 NWM only"],
        ["Planning RPE", "0.35 NWM only"],
        ["16s FVD", "200.969 NWM"],
    ]
    add_table(slide, 6.8, 1.58, 5.55, 4.2, paper_rows, font_size=13)
    add_source(slide, "Source: Navigation World Models.pdf Table 1, Table 6, Table 7.")

    # 4. Experiment matrix
    slide = blank_slide(prs)
    add_header(slide, "우리 실험 매트릭스", "Models")
    matrix_rows = [
        ["Family", "Resolution", "Condition", "Checkpoint", "Evaluated"],
        ["CDiT-S", "224", "No-Text", "nwm_cdit_s/0100000", "Frame + planning"],
        ["CDiT-S", "224", "Text Dense", "nwm_cdit_s_recon_raw_text_dense/0030000", "Frame + planning"],
        ["CDiT-S", "128", "No-Text", "nwm_cdit_s_recon_128/0030000", "Frame + planning"],
        ["CDiT-S", "128", "Text Dense", "nwm_cdit_s_recon_128_text_dense/0030000", "Frame + planning"],
        ["CDiT-B", "224", "No-Text", "nwm_cdit_b/0100000", "Frame + planning"],
        ["CDiT-B", "224", "Text Dense", "nwm_cdit_b_recon_raw_text_dense/0030000", "Frame + planning"],
        ["CDiT-B", "128", "No-Text", "nwm_cdit_b_recon_128/0030000", "Frame + planning"],
        ["CDiT-B", "128", "Text Dense", "nwm_cdit_b_recon_128_text_dense/0030000", "Frame + planning"],
    ]
    add_table(slide, 0.68, 1.48, 12.0, 4.95, matrix_rows, font_size=9)
    add_source(slide, "Source: local checkpoint and planning artifact names.")

    # 5. Evaluation setup
    slide = blank_slide(prs)
    add_header(slide, "평가 설정과 비교 범위", "Method")
    add_round_rect(slide, 0.75, 1.45, 5.7, 4.85, fill=WHITE)
    add_textbox(slide, 1.02, 1.72, 5.05, 0.32, "Future-frame prediction", size=18, color=NAVY, bold=True)
    add_bullets(
        slide,
        1.02,
        2.18,
        5.05,
        3.55,
        [
            "4초 및 16초 rollout horizon에서 LPIPS, DreamSim, FID를 측정했다.",
            "모든 지표는 낮을수록 좋다.",
            "현재 산출물에는 PSNR/FVD가 없어 논문 Table 1/6과는 부분 비교만 가능하다.",
        ],
        size=15,
    )
    add_round_rect(slide, 6.85, 1.45, 5.7, 4.85, fill=WHITE)
    add_textbox(slide, 7.12, 1.72, 5.05, 0.32, "Planning", size=18, color=TEAL, bold=True)
    add_bullets(
        slide,
        7.12,
        2.18,
        5.05,
        3.65,
        [
            "RECON 100-sample planning eval, CEM N32/K5/OPT1/rep1.",
            "ATE, RPE trans, position error, yaw error를 본다.",
            "논문 standalone planning appendix의 N=120 후보 설정과 완전히 동일한 조건은 아니다.",
        ],
        size=15,
    )
    add_source(slide, "Caveat: 논문 runtime/standalone planning 설정과 local full 100-sample planning loop는 apples-to-apples가 아니다.")

    # 6. S frame prediction table
    slide = blank_slide(prs)
    add_header(slide, "Future-frame prediction: CDiT-S", "Results")
    s_frame = frame[frame["model"] == "CDiT-S"].sort_values(["resolution", "condition"])
    add_metric_table(slide, s_frame, ["lpips4", "dream4", "fid4", "lpips16", "dream16", "fid16"], 0.62, 1.45, 7.75, 3.2)
    add_bullets(
        slide,
        8.65,
        1.55,
        3.85,
        4.5,
        [
            "128 Text는 LPIPS/DreamSim에서 4s와 16s 모두 개선된다.",
            "224 Text도 S에서는 frame metric이 전반적으로 개선된다.",
            "하지만 뒤의 planning 결과에서는 S text가 ATE를 낮추지는 못한다.",
        ],
        size=16,
    )
    add_source(slide, "Source: time-eval outputs summarized into data/frame_prediction_metrics.csv.")

    # 7. B frame prediction table
    slide = blank_slide(prs)
    add_header(slide, "Future-frame prediction: CDiT-B", "Results")
    b_frame = frame[frame["model"] == "CDiT-B"].sort_values(["resolution", "condition"])
    add_metric_table(slide, b_frame, ["lpips4", "dream4", "fid4", "lpips16", "dream16", "fid16"], 0.62, 1.45, 7.75, 3.2)
    add_bullets(
        slide,
        8.65,
        1.55,
        3.85,
        4.6,
        [
            "128 Text는 B에서도 LPIPS/DreamSim 개선이 유지된다.",
            "224 Text는 LPIPS/DreamSim이 소폭 나빠지고, 16s FID만 개선된다.",
            "B는 absolute frame metric이 S보다 낮지만, text 효과 방향은 해상도별로 다르다.",
        ],
        size=16,
    )
    add_source(slide, "Source: time-eval outputs summarized into data/frame_prediction_metrics.csv.")

    # 8. Frame delta heatmap
    slide = blank_slide(prs)
    add_header(slide, "Text 효과: frame metric delta", "Text minus No-Text")
    add_picture(slide, figures["prediction_delta"], 0.78, 1.5, width=8.5)
    add_bullets(
        slide,
        9.55,
        1.58,
        2.9,
        4.7,
        [
            "저해상도 128에서는 S/B 모두 text가 주요 perception metric을 낮춘다.",
            "224에서는 S와 B의 방향이 갈라진다.",
            "FID는 horizon과 모델 크기에 따라 덜 안정적이다.",
        ],
        size=15,
    )
    add_source(slide, "Negative percent means improvement because all listed frame metrics are lower-is-better.")

    # 9. Checkpoint curves
    slide = blank_slide(prs)
    add_header(slide, "Checkpoint-wise rollout trend", "Training/Eval Trace")
    add_picture(slide, S_TIME_CURVE, 0.7, 1.52, width=5.85)
    add_picture(slide, B_TIME_CURVE, 6.78, 1.52, width=5.85)
    add_textbox(slide, 0.9, 6.15, 5.25, 0.28, "CDiT-S checkpoint sweep", size=12, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.95, 6.15, 5.25, 0.28, "CDiT-B checkpoint sweep", size=12, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, f"Sources: {S_TIME_CURVE.relative_to(ROOT)}; {B_TIME_CURVE.relative_to(ROOT)}")

    # 10. Planning setup
    slide = blank_slide(prs)
    add_header(slide, "Planning 평가: 0427 이후 추가된 핵심 축", "Navigation")
    add_bullets(
        slide,
        0.75,
        1.48,
        5.9,
        4.9,
        [
            "0427 자료에서는 future-frame metric 중심이었다.",
            "이후 CEM planning 평가를 추가해 생성 품질이 navigation 성능으로 이어지는지 확인했다.",
            "ATE/RPE는 실제 궤적과 예측 궤적의 geometry를 본다.",
            "Position/Yaw error를 함께 보면 text가 방향 단서를 stabilizing하는지 분리해서 볼 수 있다.",
        ],
        size=17,
    )
    metric_rows = [
        ["Metric", "Interpretation"],
        ["ATE", "global trajectory mismatch"],
        ["RPE trans", "local transition error"],
        ["Position", "goal/pose position mismatch"],
        ["Yaw", "heading mismatch"],
    ]
    add_table(slide, 7.08, 1.72, 5.0, 3.35, metric_rows, font_size=12)
    add_source(slide, "Planning eval: RECON 100 samples, CEM N32/K5/OPT1/rep1.")

    # 11. S planning
    slide = blank_slide(prs)
    add_header(slide, "Planning results: CDiT-S", "Results")
    add_table(slide, 0.62, 1.43, 6.5, 2.0, table_from_planning(planning, "CDiT-S"), font_size=8)
    add_picture(slide, S_PLANNING_BARS, 0.74, 3.65, width=6.25)
    add_bullets(
        slide,
        7.52,
        1.55,
        4.82,
        4.85,
        [
            "224 No-Text가 S에서 ATE/RPE 기준 최선이다.",
            "128 Text는 yaw를 18.3% 낮추지만 ATE는 4.7% 증가한다.",
            "224 Text는 frame metric은 좋아졌지만 planning ATE/RPE/Pos/Yaw가 모두 악화된다.",
        ],
        size=16,
    )
    add_source(slide, f"Source: {S_PLANNING_CSV.relative_to(ROOT)}")

    # 12. B planning
    slide = blank_slide(prs)
    add_header(slide, "Planning results: CDiT-B", "Results")
    add_table(slide, 0.62, 1.43, 6.5, 2.0, table_from_planning(planning, "CDiT-B"), font_size=8)
    add_picture(slide, B_PLANNING_BARS, 0.74, 3.65, width=6.25)
    add_bullets(
        slide,
        7.52,
        1.55,
        4.82,
        4.85,
        [
            "128 Text는 ATE가 거의 동일하고 RPE/Pos/Yaw가 모두 개선된다.",
            "224 Text는 ATE/Pos는 악화되지만 RPE/Yaw는 개선된다.",
            "B는 S보다 yaw 개선이 일관적이며, best ATE는 224 No-Text다.",
        ],
        size=16,
    )
    add_source(slide, f"Source: {B_PLANNING_CSV.relative_to(ROOT)}")

    # 13. Direct S vs B trend answer
    slide = blank_slide(prs)
    add_header(slide, "CDiT-S와 경향성은 유사한가", "Trend")
    add_picture(slide, figures["planning_delta"], 0.72, 1.58, width=6.95)
    add_bullets(
        slide,
        8.05,
        1.55,
        4.3,
        4.95,
        [
            "큰 방향은 유사하다: text가 frame metric을 개선해도 planning ATE가 자동으로 좋아지지는 않는다.",
            "차이는 B에서 더 선명하다: B는 RPE/Yaw 개선이 더 일관적이고 128에서는 Pos도 개선된다.",
            "해상도 경향도 유사하다: 224 No-Text가 ATE 최선인 반면, text는 방향/semantic 단서에는 더 잘 작동한다.",
            "따라서 text conditioning은 semantic prior로는 유효하지만 metric geometry 보강은 별도 축으로 봐야 한다.",
        ],
        size=15,
    )
    add_source(slide, "Planning delta uses Text minus No-Text at the same model family and resolution.")

    # 14. Paper planning comparison
    slide = blank_slide(prs)
    add_header(slide, "원 논문 대비: planning", "Paper Comparison")
    add_picture(slide, figures["paper_planning"], 0.55, 1.45, width=8.4)
    add_bullets(
        slide,
        9.35,
        1.5,
        3.1,
        4.75,
        [
            "논문 NWM only: ATE 1.13, RPE 0.35.",
            "우리 best ATE는 B 224 No-Text 1.162로 paper 대비 +2.8%.",
            "우리 best RPE는 B 224 Text 0.354로 paper 대비 +1.2%.",
            "Text variant는 대체로 paper ATE까지는 못 내려간다.",
        ],
        size=15,
    )
    add_source(slide, "Source: Navigation World Models.pdf Table 7; local planning_metrics.csv.")

    # 15. Paper-style CEM budget addendum
    cem_df, cem_bullets = paper_cem_status()
    if cem_df is not None:
        slide = blank_slide(prs)
        add_header(slide, "CEM budget ablation: paper-style N120/rep3", "Planning Addendum")
        if PAPER_CEM_FIG.exists():
            add_picture(slide, PAPER_CEM_FIG, 0.62, 1.45, width=7.7)
            bullet_left = 8.65
            bullet_width = 3.85
        else:
            bullet_left = 0.8
            bullet_width = 11.4
        add_bullets(slide, bullet_left, 1.55, bullet_width, 4.75, cem_bullets, size=15)
        add_source(slide, f"Source: {PAPER_CEM_CSV.relative_to(ROOT)}; paper reference ATE 1.13 / RPE 0.35 from Table 7.")

    # 16. Paper prediction comparison
    slide = blank_slide(prs)
    add_header(slide, "원 논문 대비: 4초 future-frame metric", "Paper Comparison")
    add_picture(slide, figures["paper_prediction"], 0.55, 1.45, width=8.4)
    add_bullets(
        slide,
        9.35,
        1.5,
        3.1,
        4.75,
        [
            "논문 action+time: LPIPS 0.295, DreamSim 0.091.",
            "B 128 Text는 LPIPS 0.262로 paper reference보다 낮다.",
            "DreamSim은 모든 local variant가 paper reference보다 높다.",
            "PSNR/FVD가 아직 없어 image/video 분포 비교는 미완성이다.",
        ],
        size=15,
    )
    add_source(slide, "Source: Navigation World Models.pdf Table 1; local frame_prediction_metrics.csv.")

    # 16. Runtime
    slide = blank_slide(prs)
    add_header(slide, "Runtime과 비용", "Compute")
    add_picture(slide, figures["runtime"], 0.7, 1.5, width=7.55)
    add_bullets(
        slide,
        8.75,
        1.58,
        3.75,
        4.85,
        [
            "128 planning은 대략 0.74-1.20시간 범위다.",
            "224 planning은 설정/모델에 따라 1.00-3.17시간까지 증가한다.",
            "논문 Table 8은 단일 trajectory simulation runtime이므로 여기의 full eval total time과 직접 비교하면 안 된다.",
            "Time skip, distillation, quantization은 paper-style runtime 최적화 후보로 남아 있다.",
        ],
        size=15,
    )
    add_source(slide, "Source: local planning total_time fields; Navigation World Models.pdf Table 8.")

    # 17. Interpretation
    slide = blank_slide(prs)
    add_header(slide, "해석: text는 semantic prior, geometry는 별도 문제", "Interpretation")
    add_round_rect(slide, 0.75, 1.55, 3.7, 4.75, fill=LIGHT_TEAL)
    add_textbox(slide, 1.05, 1.9, 3.1, 0.34, "What text helps", size=18, color=TITLE, bold=True)
    add_bullets(
        slide,
        1.05,
        2.45,
        3.1,
        2.9,
        [
            "low-res visual ambiguity",
            "semantic object/scene prior",
            "heading/yaw cues in B",
        ],
        size=15,
    )
    add_round_rect(slide, 4.85, 1.55, 3.7, 4.75, fill=LIGHT_ORANGE)
    add_textbox(slide, 5.15, 1.9, 3.1, 0.34, "What text misses", size=18, color=TITLE, bold=True)
    add_bullets(
        slide,
        5.15,
        2.45,
        3.1,
        2.9,
        [
            "metric position",
            "local obstacle geometry",
            "consistent long-horizon state",
        ],
        size=15,
    )
    add_round_rect(slide, 8.95, 1.55, 3.7, 4.75, fill=WHITE)
    add_textbox(slide, 9.25, 1.9, 3.1, 0.34, "Implication", size=18, color=TITLE, bold=True)
    add_bullets(
        slide,
        9.25,
        2.45,
        3.1,
        2.9,
        [
            "better image metrics do not guarantee better ATE",
            "planning objective needs geometry-aware conditioning",
            "text should be fused with pose/RAG, not used alone",
        ],
        size=15,
    )
    add_source(slide, "Interpretation based on frame metric and planning metric divergence across S/B variants.")

    # 18. Next steps
    slide = blank_slide(prs)
    add_header(slide, "다음 실험 제안", "Next")
    add_bullets(
        slide,
        0.78,
        1.48,
        11.5,
        5.25,
        [
            "S planning을 B와 동일한 repeat sampler path로 재실행해 sampler 차이를 제거한다.",
            "PSNR/FVD를 추가해 논문 Table 1/6과 더 가까운 비교 축을 만든다.",
            "Prompt ablation: short/full/numeric/landmark caption을 분리해 text 효과의 원인을 좁힌다.",
            "Position-aware text 또는 retrieval-augmented context를 넣어 semantic prior에 metric geometry를 결합한다.",
            "Text가 yaw를 낮추지만 ATE를 높이는 route를 failure slice로 모아 rollout을 시각 점검한다.",
            "Paper runtime 최적화 경로에 맞춰 time skip, distillation, 4-bit quantization을 따로 측정한다.",
        ],
        size=18,
    )
    add_source(slide, "The highest-priority cleanup is metric-compatible comparison and sampler-aligned planning.")

    prs.save(PPTX_PATH)


def write_notes(frame: pd.DataFrame, planning: pd.DataFrame, paper_plan: pd.DataFrame) -> None:
    best_ate = planning.loc[planning["ate"].idxmin()]
    best_rpe = planning.loc[planning["rpe_trans"].idxmin()]
    b224 = planning[
        (planning["model"] == "CDiT-B")
        & (planning["resolution"] == 224)
        & (planning["condition"] == "No-Text")
    ].iloc[0]
    b128_text = frame[
        (frame["model"] == "CDiT-B")
        & (frame["resolution"] == 128)
        & (frame["condition"] == "Text")
    ].iloc[0]
    notes = f"""# NWM Text-Conditioned Update Deck Notes

Generated by `scripts/analysis/build_nwm_text_conditioned_deck.py`.

## Main takeaways

- CDiT-S와 큰 방향은 유사하다. Text는 frame metric을 개선하는 경우가 많지만 planning ATE 개선으로 자동 연결되지는 않는다.
- CDiT-B는 S보다 RPE/Yaw 개선이 더 일관적이다. 128 Text는 ATE가 거의 동일하면서 RPE/Pos/Yaw가 개선된다.
- Best local ATE: {best_ate['variant']} = {best_ate['ate']:.3f}.
- Best local RPE: {best_rpe['variant']} = {best_rpe['rpe_trans']:.3f}.
- Paper NWM only reference: ATE 1.13, RPE 0.35 from Navigation World Models Table 7.
- {b224['variant']} ATE is {(b224['ate'] / 1.13 - 1) * 100:.1f}% above the paper NWM-only ATE reference.
- {b128_text['variant']} 4s LPIPS is {(b128_text['lpips4'] / 0.295 - 1) * 100:.1f}% vs the paper action+time LPIPS reference, while 4s DreamSim is {(b128_text['dream4'] / 0.091 - 1) * 100:.1f}% vs paper.

## Data sources

- `Navigation World Models.pdf`: Table 1, Table 6, Table 7, Table 8.
- `0427 NWM Text-Conditioned 컴퓨터비전.pdf`: prior problem framing and future-work slides.
- `{S_PLANNING_CSV.relative_to(ROOT)}`.
- `{B_PLANNING_CSV.relative_to(ROOT)}`.
- `{S_TIME_CURVE.relative_to(ROOT)}`.
- `{B_TIME_CURVE.relative_to(ROOT)}`.

## Caveats

- Local prediction comparison lacks PSNR and FVD, so paper Table 1/6 comparison is partial.
- Local planning uses RECON 100-sample CEM N32/K5/OPT1/rep1. It is not identical to every paper planning/runtime setup.
- S and B planning artifacts use different path names (`CEM_N32...` vs `CEM_repeat_N32...`), so sampler alignment should be cleaned up before claiming final model-ranking conclusions.
"""
    NOTES_PATH.write_text(notes, encoding="utf-8")


def render_pdf() -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(OUT_DIR),
            str(PPTX_PATH),
        ],
        check=True,
        cwd=ROOT,
    )
    pdf_path = OUT_DIR / f"{PPTX_PATH.stem}.pdf"
    return pdf_path if pdf_path.exists() else None


def main() -> None:
    ensure_dirs()
    frame, planning, paper_plan, paper_pred = save_data()
    figures = make_figures(frame, planning, paper_plan)
    build_deck(frame, planning, paper_plan, figures)
    write_notes(frame, planning, paper_plan)
    pdf_path = render_pdf()
    print(f"Wrote {PPTX_PATH.relative_to(ROOT)}")
    if pdf_path:
        print(f"Wrote {pdf_path.relative_to(ROOT)}")
    print(f"Wrote {NOTES_PATH.relative_to(ROOT)}")
    print("Data files:")
    for path in sorted(DATA_DIR.glob("*.csv")):
        print(f"  {path.relative_to(ROOT)}")
    print("Figures:")
    for name, path in figures.items():
        print(f"  {name}: {path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
